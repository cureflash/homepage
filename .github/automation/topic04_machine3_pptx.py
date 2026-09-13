from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

OUT='denken-shinkansen/03_shinkansen_machine_3/topics/04_induction_motor/04_induction_motor_images.pptx'
prs=Presentation()
prs.slide_width=Inches(13.333333)
prs.slide_height=Inches(7.5)

FONT='Noto Sans CJK JP'
NAVY=RGBColor(20,38,62); BLUE=RGBColor(44,95,170); TEAL=RGBColor(33,140,130)
GREEN=RGBColor(61,140,85); AMBER=RGBColor(213,142,45); RED=RGBColor(183,70,70)
LIGHT=RGBColor(246,248,251); MID=RGBColor(226,231,238); DARK=RGBColor(35,40,48)
GRAY=RGBColor(95,103,114); WHITE=RGBColor(255,255,255)

def set_bg(slide, color=LIGHT):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=color

def add_title(slide, title, subtitle=None):
    box=slide.shapes.add_textbox(Inches(0.45),Inches(0.25),Inches(12.4),Inches(0.65))
    tf=box.text_frame; tf.clear(); tf.margin_left=0; tf.margin_right=0
    p=tf.paragraphs[0]; p.text=title; p.font.name=FONT; p.font.size=Pt(24); p.font.bold=True; p.font.color.rgb=NAVY
    if subtitle:
        sub=slide.shapes.add_textbox(Inches(0.48),Inches(0.88),Inches(12.0),Inches(0.38))
        stf=sub.text_frame; stf.clear(); stf.margin_left=0
        sp=stf.paragraphs[0]; sp.text=subtitle; sp.font.name=FONT; sp.font.size=Pt(10.5); sp.font.color.rgb=GRAY

def add_footer(slide, text='Topic 04 / 電験三種 機械・誘導機'):
    line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.02),Inches(7.16),Inches(13.28),Inches(0.27))
    line.fill.solid(); line.fill.fore_color.rgb=NAVY; line.line.fill.background()
    tb=slide.shapes.add_textbox(Inches(0.45),Inches(7.18),Inches(12.2),Inches(0.18))
    tf=tb.text_frame; tf.clear(); tf.margin_left=0; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=text; p.font.name=FONT; p.font.size=Pt(8.5); p.font.color.rgb=WHITE

def text_box(slide,x,y,w,h,text,size=16,bold=False,color=DARK,fill=None,line=None,align=PP_ALIGN.LEFT,margin=0.08,radius=False):
    if fill is not None:
        shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb=fill
        if line is None: shape.line.fill.background()
        else: shape.line.color.rgb=line; shape.line.width=Pt(1)
        box=shape
    else: box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin)
    p=tf.paragraphs[0]; p.text=text; p.alignment=align; p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    return box

def multi_text(slide,x,y,w,h,paras,fill=None,line=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)) if fill else slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb=fill
        if line: shape.line.color.rgb=line
        else: shape.line.fill.background()
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.12); tf.margin_top=tf.margin_bottom=Inches(0.1)
    for i,(txt,size,bold,color,level) in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=txt; p.level=level
        p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color; p.space_after=Pt(4)
    return shape

def arrow(slide,x1,y1,x2,y2,color=GRAY,width=2):
    c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(width); c.line.end_arrowhead=True

s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
add_title(s,'300系でモーターが全部変わった','Topic 04固定範囲：三相誘導電動機の構造・回転磁界・同期速度・滑り')
text_box(s,0.55,1.35,3.2,0.48,'直流主電動機 → 三相誘導電動機',18,True,WHITE,NAVY,None,PP_ALIGN.CENTER,0.05,True)
text_box(s,0.72,1.95,2.85,0.82,'直流機\n整流子・ブラシが必要',16,True,DARK,WHITE,MID,PP_ALIGN.CENTER,0.05,True)
arrow(s,3.65,2.36,4.55,2.36,AMBER,3)
text_box(s,4.62,1.95,2.9,0.82,'かご形誘導電動機\n整流子・ブラシ不要',16,True,DARK,WHITE,MID,PP_ALIGN.CENTER,0.05,True)
text_box(s,7.85,1.38,4.85,0.4,'三相交流 → 回転磁界 → 誘導起電力 → 二次電流 → トルク',15,True,NAVY,RGBColor(235,242,250),None,PP_ALIGN.CENTER,0.05,True)
outer=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(8.52),Inches(2.05),Inches(3.25),Inches(3.25)); outer.fill.solid(); outer.fill.fore_color.rgb=RGBColor(224,235,248); outer.line.color.rgb=BLUE; outer.line.width=Pt(2.5)
inner=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(9.16),Inches(2.69),Inches(1.97),Inches(1.97)); inner.fill.solid(); inner.fill.fore_color.rgb=RGBColor(237,245,238); inner.line.color.rgb=GREEN; inner.line.width=Pt(2.5)
text_box(s,8.92,2.16,2.45,0.4,'固定子（三相巻線）',13,True,BLUE,None,None,PP_ALIGN.CENTER,0)
text_box(s,9.42,3.36,1.45,0.46,'回転子\n（かご形）',12,True,GREEN,None,None,PP_ALIGN.CENTER,0)
text_box(s,8.08,5.42,4.2,0.47,'回転磁界 Ns  ＞  回転子 N',18,True,RED,RGBColor(255,244,244),None,PP_ALIGN.CENTER,0.05,True)
multi_text(s,0.62,3.12,6.9,3.48,[('電験でまず使う3式',18,True,NAVY,0),('同期速度   Ns = 120 f / p',22,True,BLUE,0),('滑り       s = (Ns − N) / Ns',22,True,TEAL,0),('二次周波数 f2 = s f1',22,True,GREEN,0),('停止時 s=1、同期速度に近づくほど s→0。Ns は回転磁界、N は回転子の実速度。',13,False,DARK,0)],fill=WHITE,line=MID)
add_footer(s)

s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
add_title(s,'速度を決める：同期速度 → 滑り → 二次周波数','R5上 問3・問4、R1 問3に直結する計算の順序')
multi_text(s,0.55,1.28,5.1,5.62,[('解法アルゴリズム',18,True,NAVY,0),('1. 周波数 f と極数 p から Ns を求める',16,True,DARK,0),('2. N または s の与え方を確認する',16,True,DARK,0),('3. s=(Ns−N)/Ns または N=(1−s)Ns',16,True,DARK,0),('4. 必要なら f2=s f1 へつなぐ',16,True,DARK,0),('5. 0<s<1、N<Ns を検算する',16,True,DARK,0),('',5,False,DARK,0),('頻出ミス',15,True,RED,0),('・p は極対数ではなく極数',14,False,DARK,0),('・Ns と N を混同しない',14,False,DARK,0),('・滑りを百分率のまま式へ入れない',14,False,DARK,0)],fill=WHITE,line=MID)
text_box(s,6.05,1.42,6.55,0.75,'50 Hz・4極なら  Ns = 120×50/4 = 1500 min⁻¹',19,True,BLUE,RGBColor(236,244,253),None,PP_ALIGN.CENTER,0.06,True)
arrow(s,9.32,2.19,9.32,2.52,GRAY,2.5)
text_box(s,6.05,2.58,6.55,0.75,'N = 1440 min⁻¹ なら  s = (1500−1440)/1500 = 0.04',18,True,TEAL,RGBColor(236,249,247),None,PP_ALIGN.CENTER,0.06,True)
arrow(s,9.32,3.35,9.32,3.68,GRAY,2.5)
text_box(s,6.05,3.74,6.55,0.75,'f₂ = s f₁ = 0.04×50 = 2 Hz',20,True,GREEN,RGBColor(240,249,241),None,PP_ALIGN.CENTER,0.06,True)
text_box(s,6.27,4.92,6.12,1.18,'意味：回転子が同期速度へ近づくほど相対速度が小さくなり、二次周波数と二次誘導起電力も小さくなる。',15,False,DARK,WHITE,MID,PP_ALIGN.LEFT,0.12,True)
text_box(s,6.25,6.28,6.15,0.42,'E₂s = s E₂0 も同じ「滑り比例」で扱う',14,True,AMBER,RGBColor(255,248,235),None,PP_ALIGN.CENTER,0.04,True)
add_footer(s)

s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
add_title(s,'電力の流れ：P₁ → P₂ → Pm → Pout','R6下 問4・R1 問3・R8上 問4で要求される損失とトルクの接続')
blocks=[(0.55,1.55,2.05,0.9,'一次入力 P₁','√3 VI cosφ',NAVY),(3.0,1.55,2.05,0.9,'二次入力 P₂','P₁−鉄損−一次銅損',BLUE),(5.45,1.55,2.05,0.9,'機械変換 Pm','(1−s)P₂',TEAL),(7.9,1.55,2.05,0.9,'軸出力 Pout','Pm−機械損',GREEN)]
for i,(x,y,w,h,t,sub,c) in enumerate(blocks):
    text_box(s,x,y,w,h,t+'\n'+sub,14,True,WHITE,c,None,PP_ALIGN.CENTER,0.06,True)
    if i<len(blocks)-1: arrow(s,x+w+0.05,y+0.45,x+w+0.38,y+0.45,GRAY,2.3)
text_box(s,2.98,2.68,2.1,0.62,'二次銅損 Pc₂ = sP₂',15,True,RED,RGBColor(255,241,241),None,PP_ALIGN.CENTER,0.05,True)
text_box(s,7.88,2.68,2.1,0.62,'η = Pout / P₁',16,True,NAVY,RGBColor(239,243,249),None,PP_ALIGN.CENTER,0.05,True)
multi_text(s,0.55,3.62,5.95,2.82,[('トルクは「どの角速度か」を確認',17,True,NAVY,0),('軸出力から：  T = Pout / ωm',21,True,GREEN,0),('ωm = 2πN/60',18,True,DARK,0),('同期ワットから： P₂ = T ωs',21,True,BLUE,0),('ωs = 2πNs/60',18,True,DARK,0)],fill=WHITE,line=MID)
multi_text(s,6.82,3.62,5.88,2.82,[('計算の見方',17,True,NAVY,0),('s = 0.04、P₂ = 100 kW のとき',15,False,DARK,0),('Pc₂ = 0.04×100 = 4 kW',18,True,RED,0),('Pm = 0.96×100 = 96 kW',18,True,TEAL,0),('機械損が無視なら Pout = 96 kW',16,True,GREEN,0),('→ 入力 > 出力、損失が正の値か検算',13,False,GRAY,0)],fill=WHITE,line=MID)
add_footer(s)

s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
add_title(s,'本試験5問への対応と範囲境界','教材へ仕様外論点を足さず、Topic 04固定範囲だけで解く')
rows=[('R8上 機械 問4','Ns・同期ワット・トルク・機械出力','Slide 1・3'),('R6下 機械 問4','三相入力・損失・二次銅損','Slide 3'),('R5上 機械 問3','回転磁界・誘導起電力・滑り','Slide 1・2'),('R5上 機械 問4','Ns・N・出力からトルク','Slide 2・3'),('R1 機械 問3','滑り・P₂・損失・効率','Slide 2・3')]
text_box(s,0.6,1.35,2.0,0.45,'公式過去問',14,True,WHITE,NAVY,None,PP_ALIGN.CENTER,0.04)
text_box(s,2.62,1.35,7.25,0.45,'要求事項',14,True,WHITE,NAVY,None,PP_ALIGN.CENTER,0.04)
text_box(s,9.89,1.35,2.75,0.45,'PowerPoint対応',14,True,WHITE,NAVY,None,PP_ALIGN.CENTER,0.04)
y=1.82
for q,req,cov in rows:
    text_box(s,0.6,y,2.0,0.68,q,13,True,DARK,WHITE,MID,PP_ALIGN.CENTER,0.04)
    text_box(s,2.62,y,7.25,0.68,req,13,False,DARK,WHITE,MID,PP_ALIGN.LEFT,0.09)
    text_box(s,9.89,y,2.75,0.68,cov,13,True,BLUE,WHITE,MID,PP_ALIGN.CENTER,0.04); y+=0.7
multi_text(s,0.6,5.58,12.05,1.17,[('このTopicでは扱わない',14,True,RED,0),('始動法 / Y-Δ始動 / 巻線形の比例推移 / インバータ・V/f・VVVF / ベクトル制御 / 回生制動 / 詳細等価回路',13,False,DARK,0),('300系との接続も「三相誘導電動機方式への転換」「整流子・ブラシ不要」の範囲に限定し、未確認の主電動機定格値・具体的制御方式は追加しない。',12,False,GRAY,0)],fill=RGBColor(255,247,247),line=RGBColor(230,190,190))
add_footer(s,'Topic 04 / EXAM_ALIGNMENT固定5問を可視化')
prs.save(OUT)
print(OUT)
