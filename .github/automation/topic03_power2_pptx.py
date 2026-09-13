from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
NAVY=RGBColor(23,50,77); BLUE=RGBColor(47,107,154); PALE=RGBColor(234,242,248); ORANGE=RGBColor(217,119,6); ORANGEP=RGBColor(255,243,224); GREEN=RGBColor(47,125,85); GREENP=RGBColor(233,245,239); RED=RGBColor(181,71,58); REDP=RGBColor(252,237,234); GRAY=RGBColor(107,114,128); DARK=RGBColor(31,41,55); LINE=RGBColor(216,222,232); WHITE=RGBColor(255,255,255); BG=RGBColor(247,249,252)
FONT='Noto Sans CJK JP'

def add_text(slide,text,x,y,w,h,size=14,bold=False,color=DARK,align=PP_ALIGN.LEFT):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=Inches(.03); tf.margin_top=tf.margin_bottom=Inches(.02); tf.vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=text; p.alignment=align
    r=p.runs[0]; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return tb

def title(slide,t,st=''):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=BG
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.16)); sh.fill.solid(); sh.fill.fore_color.rgb=NAVY; sh.line.fill.background()
    add_text(slide,t,.58,.34,12.1,.52,23,True,NAVY)
    if st:add_text(slide,st,.60,.88,12,.28,11,False,GRAY)
    add_text(slide,'電験二種・電力｜Topic 03',.55,7.05,4,.18,9,False,GRAY)

def box(slide,x,y,w,h,head,body,fill=WHITE,line=LINE,accent=NAVY,body_size=13):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line; sh.line.width=Pt(1)
    add_text(slide,head,x+.16,y+.12,w-.32,.32,14,True,accent)
    add_text(slide,body,x+.16,y+.52,w-.32,h-.65,body_size,False,DARK)

def arrow(slide,x1,y1,x2,y2,color=BLUE):
    ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=color; ln.line.width=Pt(2)
    try: ln.line.end_arrowhead=True
    except: pass

s=prs.slides.add_slide(blank); title(s,'川崎発電所は新幹線の負荷変動に追従できる？','コンバインドサイクルの熱流れ・熱収支・試験要求を一枚につなぐ')
labels=['圧縮機','燃焼器','GT','排熱回収','ST','発電機']; xs=[.65,2.65,4.65,6.65,8.65,10.65]
for i,(lab,x) in enumerate(zip(labels,xs)):
    fill=ORANGEP if lab=='排熱回収' else WHITE; line=ORANGE if lab=='排熱回収' else LINE
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.55),Inches(1.45),Inches(.9)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line
    add_text(s,lab,x+.05,1.80,1.35,.32,14,True,NAVY,PP_ALIGN.CENTER)
    if i<5: arrow(s,x+1.45,2.00,xs[i+1]-.12,2.00)
add_text(s,'燃料の熱をGTで使い、残る排熱をSTでもう一度使う → 総発電電力・熱効率を高める',.9,2.76,11.55,.42,16,True,DARK,PP_ALIGN.CENTER)
box(s,.7,3.45,3.75,2.42,'熱収支','Pg = PGT + PST\nPs = Pg − Paux\nηg = Pg / Qf\nηs = Ps / Qf\nηs = ηg(1−a)',PALE,BLUE,BLUE,14)
box(s,4.78,3.45,3.75,2.42,'一次で判定','機器と熱サイクル\n非常用電源・重要補機\n起動時の熱応力・振動・危険速度',WHITE,LINE,NAVY,13.2)
box(s,8.86,3.45,3.75,2.42,'二次で記述','原因→中間現象→出力→対策\n補機名→現象→損傷→目的\n発電端／送電端を区別して計算',GREENP,GREEN,GREEN,13.0)
add_text(s,'正式対象：R3一次問5 / H29二次問1 / R1二次問1 / H27一次問1 / H25一次問1',.72,6.35,11.9,.30,10.5,True,GRAY,PP_ALIGN.CENTER)

s=prs.slides.add_slide(blank); title(s,'起動・停止：蒸気タービンは「止めても電力が要る」','H25一次・R3一次・R1二次の共通軸は、安全停止までの物理現象')
steps=[('真空・軸封','蒸気投入前の条件'),('温度変化','急変→熱応力'),('昇速','振動を監視'),('危険速度','長時間滞在しない')]
for i,(a,b) in enumerate(steps):
    y=1.55+i*1.05
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(.88),Inches(y+.2),Inches(.34),Inches(.34)); c.fill.solid(); c.fill.fore_color.rgb=BLUE; c.line.fill.background()
    if i<3:
        ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(1.05),Inches(y+.54),Inches(1.05),Inches(y+1.15)); ln.line.color.rgb=LINE; ln.line.width=Pt(2)
    add_text(s,a,1.42,y,1.65,.35,15,True,NAVY); add_text(s,b,3.0,y,2.3,.35,12.5,False,DARK)
box(s,5.75,1.55,3.05,3.45,'タービン油ポンプ','停止後もロータは惰性回転する。\n\n軸受へ潤滑油を供給し、油膜切れ・過熱・損傷を防ぐ。',ORANGEP,ORANGE,ORANGE,13.2)
box(s,9.15,1.55,3.05,3.45,'ターニングギヤ','停止後のロータを低速回転させる。\n\n上下の温度差による不均一収縮を抑え、ロータの曲がりを防ぐ。',GREENP,GREEN,GREEN,13.2)
box(s,5.75,5.25,6.45,.86,'二次答案の型','補機名だけで終えず「停止後の物理現象 → 放置時の損傷 → 継続運転の目的」まで書く。',REDP,RED,RED,12.6)
add_text(s,'範囲境界：復水・給水設備の詳細構成そのものは追加しない。',.75,6.35,5.0,.28,10,False,GRAY)

s=prs.slides.add_slide(blank); title(s,'夏の高温で最大出力が落ちる因果','H29二次：途中を飛ばさず、原因→中間現象→出力への影響→対策')
items=[('大気温度↑',REDP,RED),('吸気密度↓',ORANGEP,ORANGE),('質量流量↓',PALE,BLUE),('GT出力↓',PALE,BLUE),('排熱回収↓',ORANGEP,ORANGE),('ST出力↓',GREENP,GREEN),('総出力↓',REDP,RED)]
x=.45
for i,(lab,fill,col) in enumerate(items):
    w=1.45
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.55),Inches(w),Inches(.86)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=col
    add_text(s,lab,x+.04,1.82,w-.08,.27,12.8,True,col,PP_ALIGN.CENTER)
    if i<6: arrow(s,x+w,1.98,x+w+.28,1.98,GRAY)
    x+=w+.35
box(s,.7,3.0,5.7,2.55,'負荷追従の時間差','GTの燃料投入・出力変化\n↓ 排ガス熱量\n↓ 蒸気発生量\n↓ ST出力\n\n蒸気側には熱容量があり、各段は瞬時に同率では変化しない。',WHITE,LINE,NAVY,13.2)
box(s,6.8,3.0,5.7,2.55,'改善策：吸気冷却','吸気を冷却\n→ 密度低下を緩和\n→ 質量流量低下を緩和\n→ GT最大出力低下を緩和\n→ コンバインド最大出力低下を緩和',GREENP,GREEN,GREEN,13.4)
add_text(s,'「暑いと効率が悪い」だけでは不足。因果の鎖を答案にする。',.75,6.05,11.8,.42,16,True,RED,PP_ALIGN.CENTER)

s=prs.slides.add_slide(blank); title(s,'川崎発電所への接続：実値と仮定値を混ぜない','確認済み実値を固定し、供給余力は教材上の仮定条件で計算する')
box(s,.65,1.42,3.55,2.18,'一次資料で確認した実値','総出力 809 MW\n発電機 4台\n燃料：都市ガス・天然ガス\n2024年度発電効率 43.1%',PALE,BLUE,BLUE,14)
box(s,.65,3.92,3.55,1.45,'注意','809 MWの全量を任意時刻に新幹線へ使える、とは意味しない。43.1%も瞬時の負荷率―効率曲線ではない。',ORANGEP,ORANGE,ORANGE,12.2)
box(s,4.55,1.42,8.05,4.62,'複合・ひっかけ例題（すべて仮定値）','夏季最大発電端 720 MW / 所内率 4.0% / 負荷 670 MW→+35 MW\n\n最大送電端：720×0.96 = 691.2 MW\n増加後負荷：670+35 = 705 MW\n必要発電端：705/0.96 = 734.375 MW\n\n734.375 > 720 → 発電端で14.375 MW不足\n705−691.2 = 13.8 MW → 送電端でも不足',WHITE,LINE,NAVY,14.0)
add_text(s,'比較する端点をそろえる：送電端負荷 ↔ 送電端能力 ／ 発電端必要出力 ↔ 発電端上限',4.82,5.55,7.5,.36,12.5,True,GREEN,PP_ALIGN.CENTER)
add_text(s,'未確認の個別効率曲線・列車1本電力・実設備ランプレートは使用していない。',.7,6.36,11.85,.28,10,False,GRAY,PP_ALIGN.CENTER)

out = Path('denken-shinkansen/06_shinkansen_power_2/topics/03_thermal_power/03_thermal_power_images.pptx')
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out))
