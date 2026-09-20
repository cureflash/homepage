from pathlib import Path
from pptx import Presentation
from PIL import Image

ppt=Path('denken-shinkansen/05_shinkansen_vehicle_2/topics/31_auxiliary_power_converter/31_auxiliary_power_converter_images.pptx')
prs=Presentation(str(ppt))
assert len(prs.slides)==6
bad=[]
for si,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+1000 or sh.top+sh.height>prs.slide_height+1000:
            bad.append((si,sh.name))
assert not bad, bad
text='\n'.join(sh.text for s in prs.slides for sh in s.shapes if hasattr(sh,'text_frame'))
required=['インバータ','コンバータ','三相負荷','力率','高調波','変換損失','効率','負荷変動','R7 一次','R5 二次','R4 一次','R2 二次','H22 一次','27 / 27']
missing=[x for x in required if x not in text]
assert not missing, missing
renders=sorted(Path('/tmp/topic31_render').glob('slide-*.png'))
assert len(renders)==6
for p in renders:
    im=Image.open(p).convert('RGB').resize((320,180))
    pix=list(im.getdata())
    nonwhite=sum(1 for r,g,b in pix if min(r,g,b)<242)
    assert nonwhite/len(pix)>0.025,(p,nonwhite/len(pix))
print('pptx/render checks PASS')
