from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
import math, os

OUT='denken-shinkansen/05_shinkansen_vehicle_2/topics/31_auxiliary_power_converter/31_auxiliary_power_converter_images.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG=RGBColor(248,250,252); NAVY=RGBColor(22,39,67); BLUE=RGBColor(49,92,166); CYAN=RGBColor(0,145,170); ORANGE=RGBColor(207,126,55); RED=RGBColor(174,69,69); GRAY=RGBColor(89,100,116); LIGHT=RGBColor(230,236,243); WHITE=RGBColor(255,255,255); BLACK=RGBColor(31,35,40); GREEN=RGBColor(55,130,105)
FONT='Noto Sans CJK JP'

def bg(s):
    f=s.background.fill; f.solid(); f.fore_color.rgb=BG

def shape(s,x,y,w,h,fill,kind=MSO_SHAPE.ROUNDED_RECTANGLE,linec=None):
    q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill
    if linec: q.line.color.rgb=linec; q.line.width=Pt(1)
    else: q.line.fill.background()
    return q

def ln(s,x1,y1,x2,y2,c=GRAY,w=1.5):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(w); return q

def tx(s,t,x,y,w,h,sz=18,c=BLACK,b=False,a=PP_ALIGN.LEFT,v=MSO_ANCHOR.TOP):
    q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=q.text_frame; f.clear(); f.word_wrap=True; f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=Inches(.04); f.vertical_anchor=v
    p=f.paragraphs[0]; p.text=t; p.alignment=a; r=p.runs[0]; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c; return q

def ttl(s,n,t,sub=''):
    tx(s,f'{n:02d}',.35,.19,.55,.45,15,CYAN,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE); tx(s,t,1.0,.16,11.9,.52,27,NAVY,True,v=MSO_ANCHOR.MIDDLE); ln(s,.35,.80,12.95,.80,NAVY,1.4)
    if sub: tx(s,sub,1.0,.68,11.7,.28,10,GRAY)

def pill(s,t,x,y,w,c):
    shape(s,x,y,w,.35,c); tx(s,t,x+.05,y+.015,w-.1,.30,10,WHITE,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

def boxtext(s,head,body,x,y,w,h,accent=BLUE):
    shape(s,x,y,w,h,WHITE,linec=LIGHT); shape(s,x,y,.08,h,accent,kind=MSO_SHAPE.RECTANGLE); tx(s,head,x+.22,y+.14,w-.35,.32,15,NAVY,True); tx(s,body,x+.22,y+.55,w-.35,h-.65,12,GRAY)

def arrow(s,x1,y1,x2,y2,c=CYAN):
    q=ln(s,x1,y1,x2,y2,c,2.1); q.line.end_arrowhead=True; return q

# 1 overview
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); ttl(s,1,'補助電源変換器 — 電験二種で解くための全体像','Topic 31 / SPEC固定8項目のみ')
steps=[('入力電源',BLUE),('コンバータ / 整流段',CYAN),('直流中間回路',ORANGE),('インバータ',BLUE),('三相補機負荷',GREEN)]
xs=[.55,2.75,5.35,7.95,10.35]
for i,(t,c) in enumerate(steps):
    shape(s,xs[i],1.25,2.0,.82,WHITE,linec=c); tx(s,t,xs[i]+.08,1.38,1.84,.52,14,c,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    if i<4: arrow(s,xs[i]+2.02,1.66,xs[i+1]-.08,1.66)
tx(s,'特定新幹線の回路・定格ではなく、試験論点を整理する一般モデル',.65,2.24,12.0,.38,14,RED,True,PP_ALIGN.CENTER)
items=['インバータ','コンバータ','三相負荷','力率','高調波','変換損失','効率','負荷変動']
for i,t in enumerate(items): pill(s,t,.65+(i%4)*3.05,2.92+(i//4)*.58,2.65,CYAN if i<4 else BLUE)
boxtext(s,'固定過去問ゲート','R7一次問4  5/5\nR5二次問3  7/7\nR4一次問4  5/5',.65,4.35,3.75,2.3,BLUE)
boxtext(s,'一次・二次を両方通す','R2二次問2  5/5\nH22一次問3  5/5\n合計 27 / 27答案要素',4.75,4.35,3.75,2.3,CYAN)
boxtext(s,'範囲境界','PV/MPPT固有事項を追加しない\n能動フィルタを実車搭載としない\n未確認実車値を真値化しない',8.85,4.35,3.75,2.3,ORANGE)

# 2 inverter / PWM
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); ttl(s,2,'電圧形インバータと誘導性負荷','R7 一次「機械」問4 — 5 / 5')
boxtext(s,'① 電流は瞬時に0にならない','誘導性負荷では、主スイッチOFF後も電流が流れ続ける経路が必要。',.55,1.15,3.7,1.2,BLUE)
boxtext(s,'② 逆並列ダイオード','電流の還流経路を与える。MOSFET等の自己消弧形素子と組み合わせる。',.55,2.55,3.7,1.35,CYAN)
boxtext(s,'③ PWM','信号波と搬送波を比較してパルス列を生成。搬送波周波数はスイッチング頻度に関係。',.55,4.10,3.7,1.45,ORANGE)
# schematic bridge
shape(s,4.7,1.25,7.9,4.9,WHITE,linec=LIGHT); tx(s,'最小回路イメージ',4.95,1.42,2.4,.35,16,NAVY,True)
for x in [5.45,7.4]:
    shape(s,x,2.15,1.25,.56,LIGHT); tx(s,'SW',x,2.20,1.25,.40,13,NAVY,True,PP_ALIGN.CENTER)
    shape(s,x,4.14,1.25,.56,LIGHT); tx(s,'SW',x,4.19,1.25,.40,13,NAVY,True,PP_ALIGN.CENTER)
    ln(s,x+.62,2.71,x+.62,4.14,NAVY,2)
ln(s,5.0,1.93,8.95,1.93,NAVY,2); ln(s,5.0,4.97,8.95,4.97,NAVY,2); tx(s,'+Vdc',4.82,1.66,.7,.3,11,RED,True); tx(s,'−',4.95,4.96,.4,.3,13,GRAY,True)
shape(s,9.35,2.64,2.35,1.52,BG,linec=GREEN); tx(s,'誘導性負荷\nL を含む',9.35,2.84,2.35,.82,18,GREEN,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE); arrow(s,8.02,3.42,9.28,3.42,GREEN)
# diode/current path
for x in [6.02,7.97]:
    tx(s,'↕ diode',x,3.12,.9,.32,10,CYAN,True,PP_ALIGN.CENTER)
arrow(s,11.75,3.85,11.75,2.90,ORANGE); tx(s,'OFF後も電流継続',9.55,4.55,2.7,.35,12,ORANGE,True,PP_ALIGN.CENTER)
# pwm strip
for i in range(18):
    if i%3!=1: shape(s,4.95+i*.40,5.53,.26,.34,CYAN,kind=MSO_SHAPE.RECTANGLE)
tx(s,'PWMパルス列 ≠ 正弦波そのもの → 基本波 + 高調波として扱う',5.0,5.98,7.1,.40,13,NAVY,True,PP_ALIGN.CENTER)
pill(s,'誘導性負荷',.65,6.34,2.1,BLUE); pill(s,'逆並列ダイオード',2.93,6.34,2.45,CYAN); pill(s,'MOSFET',5.55,6.34,1.65,BLUE); pill(s,'PWM',7.38,6.34,1.45,CYAN); pill(s,'スイッチング周波数',9.0,6.34,3.1,ORANGE)

# 3 three-phase PWM / power
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); ttl(s,3,'三相PWM・三相負荷・力率','H22 一次「機械」問3 — 5 / 5')
shape(s,.55,1.15,5.9,2.25,WHITE,linec=LIGHT); tx(s,'H22型の線間基本波実効値',.80,1.36,5.4,.35,17,NAVY,True); tx(s,'V_LL1 = √3 E_d k / (2√2)',.80,1.92,5.4,.52,25,BLUE,True,PP_ALIGN.CENTER); tx(s,'k = 2√2 V_LL1 / (√3 E_d)',.80,2.55,5.4,.40,17,CYAN,True,PP_ALIGN.CENTER); tx(s,'係数はPWM方式・電圧定義で変わる。問題文の定義を優先。',.80,3.02,5.4,.25,10,RED,True,PP_ALIGN.CENTER)
shape(s,6.8,1.15,5.9,2.25,WHITE,linec=LIGHT); tx(s,'平衡三相負荷',7.05,1.36,5.4,.35,17,NAVY,True); tx(s,'S = √3 V_LL I\nP = √3 V_LL I cosφ\nQ = √3 V_LL I sinφ',7.05,1.82,5.4,1.30,21,GREEN,True,PP_ALIGN.CENTER)
# carrier visual
shape(s,.55,3.72,6.0,2.65,WHITE,linec=LIGHT); tx(s,'三角搬送波と信号波',.80,3.90,5.5,.35,15,NAVY,True)
pts=[]
for i in range(33):
    x=.95+i*.16; tri=(i%8); yy=5.18-(tri if tri<=4 else 8-tri)*.22; pts.append((x,yy))
for a,b in zip(pts,pts[1:]): ln(s,a[0],a[1],b[0],b[1],GRAY,1)
pts2=[]
for i in range(33):
    x=.95+i*.16; yy=5.18-.52*math.sin(i/32*2*math.pi); pts2.append((x,yy))
for a,b in zip(pts2,pts2[1:]): ln(s,a[0],a[1],b[0],b[1],CYAN,2)
tx(s,'三角搬送波',1.0,5.78,1.45,.28,10,GRAY,True); tx(s,'信号波',2.6,5.78,1.2,.28,10,CYAN,True)
shape(s,6.85,3.72,5.85,2.65,WHITE,linec=LIGHT); tx(s,'力率1制御の読み方',7.10,3.90,5.3,.35,15,NAVY,True)
# phasor
ln(s,8.0,5.35,10.75,5.35,BLUE,3); arrow(s,8.0,5.35,10.70,5.35,BLUE); tx(s,'V',10.78,5.20,.35,.35,15,BLUE,True)
arrow(s,8.0,5.35,10.42,4.64,ORANGE); tx(s,'I',10.50,4.47,.35,.35,15,ORANGE,True); tx(s,'φ',8.98,4.72,.35,.35,14,NAVY,True); tx(s,'基本波電圧・電流の位相関係から\n有効電力と力率を判断する',7.55,5.78,4.5,.50,12,GRAY,True,PP_ALIGN.CENTER)

# 4 harmonics
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); ttl(s,4,'高調波を「基本波 + 残差」に分ける','R5 二次「機械・制御」問3 — 7 / 7')
shape(s,.55,1.16,4.15,5.35,WHITE,linec=LIGHT); tx(s,'分解と補償',.82,1.37,3.6,.35,17,NAVY,True); tx(s,'i_L(t) = i_1(t) + i_H(t)',.82,2.02,3.6,.42,20,BLUE,True,PP_ALIGN.CENTER); tx(s,'i_c(t) = −i_H(t)',.82,2.72,3.6,.42,20,CYAN,True,PP_ALIGN.CENTER); tx(s,'i_s(t) = i_1(t)',.82,3.42,3.6,.42,20,GREEN,True,PP_ALIGN.CENTER); arrow(s,2.62,3.90,2.62,4.58,GREEN); tx(s,'電源側から高調波成分を打ち消す\nという試験上の補償原理',.95,4.72,3.34,.75,13,GRAY,True,PP_ALIGN.CENTER); tx(s,'※ 実車への能動フィルタ搭載を\n主張しているわけではない',.95,5.63,3.34,.55,11,RED,True,PP_ALIGN.CENTER)
shape(s,4.95,1.16,4.0,5.35,WHITE,linec=LIGHT); tx(s,'対称方形波の実効値',5.22,1.37,3.45,.35,17,NAVY,True); tx(s,'I₁ = (2√2/π) I_L\n= 約 0.900 I_L',5.22,2.01,3.45,.9,20,BLUE,True,PP_ALIGN.CENTER); tx(s,'I_H = √(I_L² − I₁²)\n= 約 0.435 I_L',5.22,3.20,3.45,.9,20,CYAN,True,PP_ALIGN.CENTER); tx(s,'I_L² = I₁² + I_H²',5.22,4.52,3.45,.42,18,GREEN,True,PP_ALIGN.CENTER); tx(s,'波形・基本波・高調波合成を\n別々の量として追う',5.32,5.42,3.25,.55,12,GRAY,True,PP_ALIGN.CENTER)
shape(s,9.2,1.16,3.5,5.35,WHITE,linec=LIGHT); tx(s,'総合力率',9.46,1.37,2.95,.35,17,NAVY,True); tx(s,'λ = (I₁/I) cosφ₁',9.46,2.03,2.95,.42,20,ORANGE,True,PP_ALIGN.CENTER); tx(s,'= cosφ₁ / √(1+THD_I²)',9.46,2.71,2.95,.42,16,ORANGE,True,PP_ALIGN.CENTER); tx(s,'成立条件',9.46,3.62,2.95,.3,14,RED,True,PP_ALIGN.CENTER); tx(s,'・電圧は正弦波\n・電流ひずみだけを考える\n・基本波と高調波が直交',9.55,4.05,2.75,1.2,13,GRAY); tx(s,'無条件の一般式にしない',9.46,5.55,2.95,.38,12,RED,True,PP_ALIGN.CENTER)

# 5 losses efficiency
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); ttl(s,5,'変換損失と効率 — 何が負荷で変わるか','R4 一次「機械」問4 5/5 + R2 二次「機械・制御」問2 5/5')
boxtext(s,'導通損失','定電圧降下モデル:\nP_cond ≈ V_on I_avg,on\n\nMOSFET抵抗モデル:\nP_cond ≈ I_rms,on² R_on',.55,1.18,3.7,2.18,BLUE)
boxtext(s,'スイッチング損失','P_sw = f_s(E_on+E_off)\n周期が短いほど単位時間当たり損失は増える。\nソフトスイッチングは切替時の重なりを減らす。',.55,3.55,3.7,2.40,CYAN)
shape(s,4.55,1.18,4.05,4.77,WHITE,linec=LIGHT); tx(s,'効率の一般式',4.82,1.40,3.50,.35,17,NAVY,True); tx(s,'η = P_out / (P_out + P_loss)',4.82,1.97,3.50,.48,20,GREEN,True,PP_ALIGN.CENTER); tx(s,'教材用モデル',4.82,2.78,3.50,.30,14,NAVY,True,PP_ALIGN.CENTER); tx(s,'P_out = x P_N\nP_loss = P₀ + x² P_vN',4.82,3.25,3.50,.80,19,BLUE,True,PP_ALIGN.CENTER); tx(s,'η(x) = xP_N /\n(xP_N+P₀+x²P_vN)',4.82,4.27,3.50,.82,17,CYAN,True,PP_ALIGN.CENTER); tx(s,'P₀=x²P_vN はこのモデル内の最大効率条件。\n変換器一般の普遍則ではない。',4.82,5.22,3.50,.55,10,RED,True,PP_ALIGN.CENTER)
# small curve
shape(s,8.9,1.18,3.8,4.77,WHITE,linec=LIGHT); tx(s,'負荷率と効率（模式）',9.15,1.40,3.3,.35,16,NAVY,True); ln(s,9.35,5.15,12.25,5.15,GRAY,1.2); ln(s,9.35,5.15,9.35,2.15,GRAY,1.2); tx(s,'負荷率 x',10.35,5.28,1.2,.3,10,GRAY); tx(s,'η',9.02,2.02,.3,.3,12,GRAY,True)
pts=[]
for i in range(21):
    x=.05+i*.0475; eta=x/(x+.02+.02*x*x) if x>0 else 0; pts.append((9.35+2.9*x,5.15-2.65*eta))
for a,b in zip(pts,pts[1:]): ln(s,a[0],a[1],b[0],b[1],ORANGE,2.3)
tx(s,'固定損と負荷依存損を分けて考える',9.25,5.62,3.15,.35,11,ORANGE,True,PP_ALIGN.CENTER)
pill(s,'オン損失',.75,6.35,1.8,BLUE); pill(s,'漏れ電流',2.72,6.35,1.8,CYAN); pill(s,'スイッチング損失',4.70,6.35,2.50,ORANGE); pill(s,'効率',7.40,6.35,1.45,GREEN); pill(s,'負荷変動',9.05,6.35,2.00,BLUE)

# 6 DC link and gate coverage
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); ttl(s,6,'負荷ステップを直流中間回路のエネルギーで読む','最終確認 — 固定5問・27答案要素を可視化へ接続')
shape(s,.55,1.15,6.0,3.05,WHITE,linec=LIGHT); tx(s,'直流中間コンデンサ',.82,1.38,5.45,.35,17,NAVY,True); tx(s,'W_dc = 1/2 C V_dc²',.82,1.98,5.45,.45,22,BLUE,True,PP_ALIGN.CENTER); tx(s,'dW_dc/dt = P_in − P_out − P_loss',.82,2.70,5.45,.45,18,CYAN,True,PP_ALIGN.CENTER); tx(s,'負荷急増直後: P_in < P_out + P_loss\n→ W_dc低下 → V_dc低下\n入力側制御が追従すると回復方向へ',.95,3.30,5.15,.72,13,GRAY,True,PP_ALIGN.CENTER)
shape(s,6.85,1.15,5.85,3.05,WHITE,linec=LIGHT); tx(s,'Vdc の模式応答',7.10,1.38,5.3,.35,17,NAVY,True); ln(s,7.45,3.55,12.15,3.55,GRAY,1.2); ln(s,7.45,3.55,7.45,1.95,GRAY,1.2); tx(s,'時間',10.95,3.66,.7,.25,10,GRAY); tx(s,'Vdc',7.05,1.84,.45,.25,10,GRAY)
pts=[(7.50,2.28),(8.35,2.28),(8.55,3.05),(9.20,3.10),(10.15,2.68),(11.10,2.42),(12.00,2.31)]
for a,b in zip(pts,pts[1:]): ln(s,a[0],a[1],b[0],b[1],ORANGE,2.5)
ln(s,8.45,1.95,8.45,3.55,RED,1.1); tx(s,'負荷↑',8.15,1.72,.8,.27,10,RED,True); tx(s,'低下',8.60,3.16,.7,.26,10,ORANGE,True); tx(s,'追従・回復',10.25,2.90,1.3,.27,10,GREEN,True)
# coverage table
shape(s,.55,4.50,12.15,1.78,WHITE,linec=LIGHT); tx(s,'過去問品質ゲート',.78,4.68,2.4,.32,16,NAVY,True)
rows=[('R7 一次 問4','5 / 5'),('R5 二次 問3','7 / 7'),('R4 一次 問4','5 / 5'),('R2 二次 問2','5 / 5'),('H22 一次 問3','5 / 5')]
for i,(a,b) in enumerate(rows):
    x=3.20+i*1.78; tx(s,a,x,4.68,1.62,.38,10,GRAY,True,PP_ALIGN.CENTER); tx(s,b,x,5.12,1.62,.40,16,BLUE if '一次' in a else CYAN,True,PP_ALIGN.CENTER)
tx(s,'一次 15 / 15     二次 12 / 12     合計 27 / 27 connected',3.18,5.72,9.18,.35,14,GREEN,True,PP_ALIGN.CENTER)
tx(s,'頻出ミス: 線間/相・基本波/実効値・力率/効率・損失エネルギー/損失電力を混同しない',.72,6.52,11.9,.40,12,RED,True,PP_ALIGN.CENTER)

# geometry guard
for i,s in enumerate(prs.slides,1):
    for q in s.shapes:
        if q.left < 0 or q.top < 0 or q.left+q.width > prs.slide_width+1000 or q.top+q.height > prs.slide_height+1000:
            raise RuntimeError(f'out of bounds slide {i}: {q.name}')
os.makedirs(os.path.dirname(OUT),exist_ok=True); prs.save(OUT)
print(OUT)
