from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = 'denken-shinkansen/07_shinkansen_machine_2/topics/22_integrated_route_simulation/22_integrated_route_simulation_images.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
FONT = 'Noto Sans CJK JP'
NAVY='10253F'; CYAN='1B91B7'; ORANGE='E58B34'; GREEN='3C9D67'; RED='C95757'; PURPLE='6A5ACD'
BG='F5F7FA'; WHITE='FFFFFF'; TEXT='1E293B'; MUTED='5B6778'; GRID='D8E0EA'; PALE='EAF0F6'

def rgb(h): return RGBColor.from_string(h)
def box(slide,x,y,w,h,text='',fill=WHITE,line=GRID,size=14,color=TEXT,bold=False,align=PP_ALIGN.CENTER,round=True):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.color.rgb=rgb(line); sh.line.width=Pt(1)
    tf=sh.text_frame; tf.clear(); tf.margin_left=tf.margin_right=Inches(.08); tf.margin_top=tf.margin_bottom=Inches(.04); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=rgb(color); r.font.bold=bold
    return sh
def tb(slide,x,y,w,h,text,size=12,color=TEXT,bold=False,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.margin_left=tf.margin_right=Inches(.03); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=rgb(color); r.font.bold=bold
    return sh
def title(slide,n,t,sub):
    box(slide,.45,.18,.85,.42,f'{n:02d}',NAVY,NAVY,20,WHITE,True)
    tb(slide,1.52,.13,11.35,.5,t,24,NAVY,True); tb(slide,1.52,.60,11.35,.28,sub,9,MUTED)
def arr(slide,x1,y1,x2,y2,c=NAVY,w=2):
    l=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); l.line.color.rgb=rgb(c); l.line.width=Pt(w); l.line.end_arrowhead=True; return l
def footer(slide,src): tb(slide,.48,7.12,12.3,.20,f'出典・条件: {src} ｜ 未確認実車値は真値化しない',8,MUTED)

s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
title(s,1,'主回路を「電力の一本道」として読む','Topic 22｜変圧器・PWM・DCリンク・VVVF・誘導機・回生を一つの符号規約で接続')
nodes=[('架線\n25 kV',.55,CYAN),('主変圧器',1.85,CYAN),('PWM',3.15,ORANGE),('DC\nリンク',4.35,ORANGE),('VVVF',5.55,ORANGE),('誘導\n電動機',6.75,PURPLE),('歯車',8.05,PURPLE),('車輪',9.15,PURPLE),('列車',10.25,GREEN),('速度\nFB',11.35,GREEN)]
for i,(lab,x,c) in enumerate(nodes):
    box(s,x,1.15,1.0,.66,lab,c,c,12,WHITE,True)
    if i<len(nodes)-1: arr(s,x+1.0,1.48,nodes[i+1][1]-.05,1.48)
box(s,.55,2.15,3.8,3.55,'',WHITE,GRID); tb(s,.82,2.37,3.25,.34,'主変圧器',18,CYAN,True,PP_ALIGN.CENTER)
tb(s,.82,2.92,3.25,1.55,'P_loss = P₀ + x²P_cu,N\nη = P_out / (P_out + P_loss)\n最大効率：P₀ = x²P_cu,N',16,TEXT,True,PP_ALIGN.CENTER)
tb(s,.82,4.66,3.25,.60,'kVAとkWを力率なしで同一視しない。\n回生でも損失は正の散逸量。',11,MUTED,False,PP_ALIGN.CENTER)
box(s,4.70,2.15,3.85,3.55,'',WHITE,GRID); tb(s,4.96,2.37,3.33,.34,'PWM / DCリンク / VVVF',18,ORANGE,True,PP_ALIGN.CENTER)
tb(s,4.98,2.90,3.28,1.62,'PWMコンバータ：交流 ↔ DCリンク\nVVVF：DCリンク ↔ 三相可変交流\n信号波 ≠ キャリア ≠ 基本波周波数',14,TEXT,True,PP_ALIGN.CENTER)
tb(s,5.02,4.68,3.20,.55,'未指定の変換器損失は追加せず、\n基準モデルでは理想変換器とする。',11,MUTED,False,PP_ALIGN.CENTER)
box(s,8.90,2.15,3.85,3.55,'',WHITE,GRID); tb(s,9.18,2.37,3.30,.34,'力行／回生の符号',18,NAVY,True,PP_ALIGN.CENTER)
box(s,9.30,3.00,2.95,.58,'力行：P > 0 / T > 0',GREEN,GREEN,14,WHITE,True); box(s,9.30,3.82,2.95,.58,'回生：P < 0 / T < 0',RED,RED,14,WHITE,True)
tb(s,9.14,4.66,3.35,.68,'P_cat = P_sec + P_loss,tr\n消費=max(P_cat,0)／回生=max(−P_cat,0)',12,TEXT,True,PP_ALIGN.CENTER)
box(s,.72,6.08,11.9,.48,'R08一次 問3 (1)〜(3) ｜ R07一次 問4 (4),(5) ｜ R06一次 問2 (1),(2)',PALE,'C9DFEA',11,NAVY,True)
footer(s,'22_integrated_route_simulation.md / SPEC.md')

s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
title(s,2,'誘導機から列車加速度まで','同期速度 → 滑り → 電力・トルク → 歯車・車輪 → 運動方程式')
chain=[('nₛ=120f/p',CYAN),('s=(nₛ−n)/nₛ',PURPLE),('P₂=3|I₂′|²r₂′/s',PURPLE),('T=P₂/ωₛ',PURPLE),('v=πDn/(60i)',GREEN),('F=2iT/D',GREEN),('a=(F−F_R)/M',GREEN)]
x=.45
for eq,c in chain: box(s,x,1.20,1.63,.62,eq,c,c,12,WHITE,True); x+=1.80
box(s,.55,2.25,3.75,3.42,'',WHITE,GRID); tb(s,.82,2.50,3.20,.36,'滑りの意味',18,PURPLE,True,PP_ALIGN.CENTER)
tb(s,.82,3.10,3.20,1.48,'0 < s < 1：電動機運転\ns = 0：同期速度\ns < 0：発電機運転 → 回生側',15,TEXT,True,PP_ALIGN.CENTER)
tb(s,.82,4.82,3.20,.54,'負の滑りを消さず、\n電力方向の符号として保持する。',11,MUTED,False,PP_ALIGN.CENTER)
box(s,4.62,2.25,3.80,3.42,'','F2FAF6','C9E6D5'); tb(s,4.92,2.49,3.20,.36,'t=120 s（力行）',18,GREEN,True,PP_ALIGN.CENTER)
tb(s,4.98,3.10,3.05,1.86,'f=112.204561 Hz\nv=34.266767 m/s\ns=+0.152235\nT=+16.874 kN·m\na=+0.277014 m/s²',14,TEXT,True,PP_ALIGN.CENTER)
box(s,8.72,2.25,4.03,3.42,'','FFF4F3','E9C8C4'); tb(s,9.05,2.49,3.35,.36,'t=210 s（回生）',18,RED,True,PP_ALIGN.CENTER)
tb(s,9.10,3.10,3.20,1.86,'f=83.594231 Hz\nv=33.307446 m/s\ns=−0.106058\nT=−13.138 kN·m\na=−0.304617 m/s²',14,TEXT,True,PP_ALIGN.CENTER)
box(s,.72,6.08,11.9,.48,'R07一次 問2 (1) ｜ R06一次 問2 (1),(2) ｜ R07二次 機械・制御 問2 (1)','F0ECFA','D8D0F2',11,NAVY,True)
footer(s,'22_integrated_route_simulation.md / calculation QA')

s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
title(s,3,'速度フィードバックと3301点シミュレーション','PI/PID → 周波数指令 → 誘導機 → 列車速度を時系列で戻す')
labs=[('r\n速度指令',.55,CYAN),('Σ\ne=r−y',2.25,NAVY),('C(s)\nPI/PID',3.95,ORANGE),('G(s)\n主回路＋列車',5.90,PURPLE),('y\n列車速度',8.20,GREEN)]
for lab,x,c in labs: box(s,x,1.15,1.40,.66,lab,c,c,12,WHITE,True)
for a,b in zip(labs[:-1],labs[1:]): arr(s,a[1]+1.40,1.48,b[1]-.05,1.48)
arr(s,9.60,1.48,10.50,1.48,GREEN,2); tb(s,10.55,1.22,2.05,.50,'単位負帰還\nE/R=1/(1+CG)',11,GREEN,True,PP_ALIGN.CENTER)
box(s,.55,2.20,4.00,2.25,'',WHITE,GRID); tb(s,.82,2.42,3.45,.32,'二次記述の式と手順',17,ORANGE,True,PP_ALIGN.CENTER)
tb(s,.85,2.95,3.36,1.18,'C(s)=Kp+Ki/s(+Kd·s)\nY/R=CG/(1+CG)\n逆ラプラス → y(t) → t=t₀ へ代入',13,TEXT,True,PP_ALIGN.CENTER)
box(s,4.78,2.20,3.80,2.25,'','FFF9EE','E9D7AE'); tb(s,5.03,2.42,3.30,.32,'教材用速度指令',17,ORANGE,True,PP_ALIGN.CENTER)
tb(s,5.10,2.92,3.15,1.25,'0–120 s：0→40 m/s\n120–180 s：40 m/s\n180–300 s：40→0 m/s\n300–330 s：0 m/s',12,TEXT,True,PP_ALIGN.CENTER)
box(s,8.82,2.20,3.93,2.25,'','EFF7FB','C9DFEA'); tb(s,9.08,2.42,3.40,.32,'計算QA',17,CYAN,True,PP_ALIGN.CENTER)
tb(s,9.08,2.92,3.40,1.25,'3301 rows／Δt=0.1 s\n力行 1828点／回生 1149点\n消費・回生 同時正値 0\n速度負値 0',12,TEXT,True,PP_ALIGN.CENTER)
for i,(v,lbl) in enumerate([('39.367','max speed [m/s]'),('4383.698','max consumption [kW]'),('2209.887','max regeneration [kW]')]):
    x=.70+i*4.15; box(s,x,4.78,3.78,.82,'',WHITE,GRID); tb(s,x+.15,4.91,1.35,.36,v,18,NAVY,True); tb(s,x+1.55,4.93,2.0,.32,lbl,10,MUTED,True)
box(s,.72,5.92,11.9,.48,'R07二次 問4 (1)〜(5)：PI・偏差・閉ループ・定常偏差・インパルス／時刻応答','F0ECFA','D8D0F2',11,NAVY,True)
footer(s,'22_integrated_route_simulation_calculation_qa.md / results.csv / six SVGs')

s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
title(s,4,'固定EXAM_ALIGNMENT 16/16 と出典','一次4問＋二次2問を維持し、仕様外の実車値・制御方式を追加しない')
rows=[('R08一次 機械 問3 (1)〜(3)','3','Slide 1：変圧器'),('R07一次 機械 問2 (1)','1','Slide 2：誘導機'),('R07一次 機械 問4 (4),(5)','2','Slide 1：PWM'),('R06一次 機械 問2 (1),(2)','2','Slide 1–2：回生'),('R07二次 機械・制御 問2 (1)','2','Slide 2：同期速度・トルク'),('R07二次 機械・制御 問4 (1)〜(5)','6','Slide 3：PI・応答')]
box(s,.55,1.15,4.55,.42,'固定過去問',NAVY,NAVY,11,WHITE,True); box(s,5.10,1.15,1.25,.42,'要素',NAVY,NAVY,11,WHITE,True); box(s,6.35,1.15,6.40,.42,'対応',NAVY,NAVY,11,WHITE,True)
y=1.57
for i,(a,b,c) in enumerate(rows):
    f='FFFFFF' if i%2==0 else 'F1F4F8'; box(s,.55,y,4.55,.46,a,f,GRID,10,TEXT,False,PP_ALIGN.LEFT,False); box(s,5.10,y,1.25,.46,b,f,GRID,10,TEXT,False,PP_ALIGN.CENTER,False); box(s,6.35,y,6.40,.46,c,f,GRID,10,TEXT,False,PP_ALIGN.LEFT,False); y+=.46
box(s,.55,4.55,3.45,.58,'16 / 16 COVERED',GREEN,GREEN,16,WHITE,True)
box(s,4.30,4.42,4.05,1.60,'',WHITE,GRID); tb(s,4.55,4.60,3.55,.28,'外部参照',16,NAVY,True,PP_ALIGN.CENTER); tb(s,4.58,5.03,3.50,.70,'電気技術者試験センター 二種 過去問・解答\ne-sysnet（誘導機／制御／PID）\n電験王2（R07一次・二次）',10,TEXT,False,PP_ALIGN.CENTER)
box(s,8.65,4.42,4.10,1.60,'',WHITE,GRID); tb(s,8.90,4.60,3.60,.28,'GitHub正本',16,NAVY,True,PP_ALIGN.CENTER); tb(s,8.92,5.03,3.56,.70,'22_integrated_route_simulation.md\ncalculation_qa.md / results.csv / six SVGs\nMASTER_SPEC / EXAM_ALIGNMENT_SPEC / series SPEC',9,TEXT,False,PP_ALIGN.CENTER)
box(s,.72,6.28,11.9,.48,'外部画像転載なし。固定過去問の個別正答・完成解答は保存せず、独立再解答は全成果物完成後に実施。','FFF4F3','E9C8C4',10,RED,True)
footer(s,'MASTER_SPEC.md / EXAM_ALIGNMENT_SPEC.md / series SPEC.md / Topic 22 source')

prs.save(OUT)
print(OUT)
