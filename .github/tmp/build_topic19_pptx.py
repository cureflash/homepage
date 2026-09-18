from pptx import Presentation
from pptx.util import Inches as I,Pt
from pptx.enum.shapes import MSO_SHAPE as S,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor as C
from scipy import signal
import numpy as np,sys
O=sys.argv[1]
F='Noto Sans CJK JP'; ink=C(28,36,48); muted=C(90,100,115); blue=C(30,90,150); green=C(20,125,105); orange=C(170,95,20); bg=C(247,248,250); white=C(255,255,255); edge=C(210,218,226); soft=C(233,240,248)
r=Presentation(); r.slide_width=I(13.333); r.slide_height=I(7.5); blank=r.slide_layouts[6]
def txt(s,t,x,y,w,h,z=14,c=ink,b=0,a=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); p=q.text_frame.paragraphs[0]; p.text=t; p.font.name=F;p.font.size=Pt(z);p.font.color.rgb=c;p.font.bold=bool(b);p.alignment=a;return q
def box(s,x,y,w,h,fill=white):
 q=s.shapes.add_shape(S.ROUNDED_RECTANGLE,I(x),I(y),I(w),I(h));q.fill.solid();q.fill.fore_color.rgb=fill;q.line.color.rgb=edge;return q
def base(n,title,sub=''):
 s=r.slides.add_slide(blank);q=s.shapes.add_shape(S.RECTANGLE,0,0,r.slide_width,r.slide_height);q.fill.solid();q.fill.fore_color.rgb=bg;q.line.fill.background();q=s.shapes.add_shape(S.RECTANGLE,0,0,r.slide_width,I(.07));q.fill.solid();q.fill.fore_color.rgb=blue;q.line.fill.background();txt(s,title,.58,.30,12.1,.55,26,ink,1);txt(s,sub,.60,.90,12,.30,10.5,muted);txt(s,'電験二種 機械・制御 / Topic 19',.52,7.14,4,.18,8,muted);txt(s,str(n),12.3,7.12,.45,.18,8,muted,0,PP_ALIGN.RIGHT);return s
def formula(s,t,x,y,w,h,c=soft):
 box(s,x,y,w,h,c);txt(s,t,x+.06,y+.08,w-.12,h-.12,13,ink,1,PP_ALIGN.CENTER)
s=base(1,'P・I・Dは「偏差のどこを見るか」が違う','定数項・1/s項・s項を見分けると補償器を判別できる')
for i,(h,a,b,d,col) in enumerate([('P：現在の偏差','uP=Kp e','C_P=Kp','比例',blue),('I：偏差の累積','uI=Ki ∫e dt','C_I=Ki/s','定常偏差へ効く',green),('D：偏差の変化','uD=Kd de/dt','C_D=Kd s','過渡へ効く',orange)]):
 x=.62+i*4.17;box(s,x,1.40,3.72,2.18);txt(s,h,x+.16,1.58,3.3,.30,15,col,1);formula(s,a,x+.18,2.04,3.34,.48);formula(s,b,x+.18,2.63,3.34,.48);txt(s,d,x+.2,3.24,3.3,.23,10,muted,0,PP_ALIGN.CENTER)
box(s,.62,3.88,12.06,2.64);txt(s,'補償器の判別',.88,4.08,2,.30,15,ink,1)
for j,(n,e,a) in enumerate([('P','Kp','比例'),('PI','Kp+Ki/s','比例＋積分'),('PD','Kp+Kd s','比例＋微分'),('PID','Kp+Ki/s+Kd s','比例＋積分＋微分')]):
 y=4.52+j*.44;txt(s,n,.96,y,.65,.25,11,blue,1);txt(s,e,1.72,y,3.25,.25,11,ink,1);txt(s,a,5.08,y,2.1,.25,10,muted)
txt(s,'時間定数表示',7.75,4.08,1.8,.30,13,ink,1);formula(s,'C(s)=Kp(1+1/(Ti s)+Td s)',7.72,4.48,4.40,.52,C(232,244,241));txt(s,'Ki=Kp/Ti   Kd=Kp·Td',7.95,5.18,3.95,.24,11,ink,1,PP_ALIGN.CENTER);txt(s,'Ti=Kp/Ki   Td=Kd/Kp',7.95,5.52,3.95,.24,11,ink,1,PP_ALIGN.CENTER);txt(s,'Kp：比例ゲイン / Ti：積分時間 / Td：微分時間',7.76,5.96,4.35,.24,9.5,muted,0,PP_ALIGN.CENTER)
s=base(2,'定常偏差は E(s) を作って最終値の定理で判定する','出力Y(s)ではなく偏差E(s)の最終値を求める')
box(s,.62,1.36,5.68,4.95);txt(s,'単位負帰還の基本式',.90,1.60,2.4,.30,15,ink,1);formula(s,'E(s)=R(s)/{1+C(s)G(s)}',.92,2.04,5.02,.58);formula(s,'e(∞)=lim[s→0] sE(s)',.92,2.80,5.02,.58,C(232,244,241));txt(s,'単位ステップ：R(s)=1/s',1.00,3.72,4.7,.26,11);txt(s,'単位ランプ：R(s)=1/s²',1.00,4.10,4.7,.26,11);txt(s,'最終値の定理は応答が収束する条件で使う。\n一般の安定判別はTopic 20へ送る。',1.00,4.72,4.8,.75,10,muted)
box(s,6.58,1.36,6.10,4.95);txt(s,'教材用プラント  G(s)=1/(2s+1)',6.88,1.60,5.0,.30,14,ink,1);txt(s,'P制御',6.92,2.16,1,.25,13,blue,1);formula(s,'C(s)=Kp',7.88,2.06,1.70,.46);txt(s,'ステップ：e(∞)=1/(1+Kp)',6.94,2.68,5.0,.27,11);txt(s,'閉ループ時定数：Tcl=2/(1+Kp)',6.94,3.02,5.0,.27,11);txt(s,'PI制御',6.92,3.72,1,.25,13,green,1);formula(s,'C(s)=Kp+Ki/s',7.88,3.61,2.60,.46,C(232,244,241));txt(s,'ステップ：e(∞)=0（収束する教材モデル）',6.94,4.22,5.2,.27,11);txt(s,'ランプ：e(∞)=1/Ki',6.94,4.58,5.0,.27,11);txt(s,'積分項があれば全入力で偏差0、とは限らない。',6.94,5.16,5.0,.30,9.5,muted)
txt(s,'接続：R07二次（PI判別・E(s)・ステップ/ランプ偏差） / R02二次（P制御・定常偏差）',.8,6.55,11.7,.28,9.5,muted,0,PP_ALIGN.CENTER)
def graph(s,x,title,series):
 box(s,x,1.52,3.76,3.15);txt(s,title,x+.18,1.72,3.4,.25,12,ink,1,PP_ALIGN.CENTER); gx=x+.50;gy=2.20;gw=2.88;gh=1.92
 for x1,y1,x2,y2 in [(gx,gy+gh,gx+gw,gy+gh),(gx,gy,gx,gy+gh)]:
  q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2));q.line.color.rgb=ink;q.line.width=Pt(.8)
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(gx),I(gy+gh*(1-1/1.6)),I(gx+gw),I(gy+gh*(1-1/1.6)));q.line.color.rgb=muted;q.line.dash_style=1;q.line.width=Pt(.6)
 cols=[blue,orange,green]
 for k,(lab,Kp,Ki,Kd) in enumerate(series):
  t=np.linspace(0,15,55);_,y=signal.step(signal.TransferFunction([Kd,Kp,Ki],[2+Kd,1+Kp,Ki]),T=t); y=np.clip(y,0,1.6)
  for i in range(len(t)-1):
   X1=gx+gw*t[i]/15;X2=gx+gw*t[i+1]/15;Y1=gy+gh*(1-y[i]/1.6);Y2=gy+gh*(1-y[i+1]/1.6);q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(X1),I(Y1),I(X2),I(Y2));q.line.color.rgb=cols[k];q.line.width=Pt(1.2)
  txt(s,lab,x+.40+k*1.05,4.28,1.0,.20,7.7,cols[k],1,PP_ALIGN.CENTER)
 txt(s,'0',gx-.08,gy+gh+.03,.20,.15,7,muted);txt(s,'15 s',gx+gw-.20,gy+gh+.03,.38,.15,7,muted);txt(s,'1.0',gx-.40,gy+gh*(1-1/1.6)-.08,.35,.16,7,muted)
s=base(3,'Kp・Ki・Kdは「1つずつ」変えて速度応答を比較する','教材用仮定モデル：G(s)=1/(2s+1)、単位負帰還・単位ステップ。実車PID値ではない')
graph(s,.48,'Kpだけを変更',[('Kp=1',1,1,.5),('Kp=2',2,1,.5),('Kp=4',4,1,.5)]);graph(s,4.55,'Kiだけを変更',[('Ki=.5',2,.5,.5),('Ki=1',2,1,.5),('Ki=2',2,2,.5)]);graph(s,8.62,'Kdだけを変更',[('Kd=0',2,1,0),('Kd=.5',2,1,.5),('Kd=1',2,1,1)])
box(s,.65,4.93,12.0,1.50);txt(s,'閉ループ分母',.93,5.13,1.45,.25,12,ink,1);formula(s,'(T+Kd)s²+(1+Kp)s+Ki',2.24,5.03,3.55,.46);txt(s,'ωn=√{Ki/(T+Kd)}    ζ=(1+Kp)/{2√[Ki(T+Kd)]}',6.08,5.12,5.9,.27,10.5,ink,1);txt(s,'読む指標：最終値 / 定常偏差 / 10–90%立上り / 最大オーバーシュート / 2%整定時間',.92,5.76,11.3,.27,10);txt(s,'係数変更の方向を暗記で一般化せず、同一条件の実際の応答を比較する。',.92,6.12,11.3,.25,9.5,muted)
s=base(4,'本試験では「式を読む → E(s) → 最終値 → 応答」の順で処理する','固定一次1問＋二次4問、計5問・12答案要素をこの章だけで接続する')
box(s,.60,1.42,5.10,5.08);txt(s,'解法フロー',.90,1.64,1.7,.30,15,ink,1)
for i,t in enumerate(['① 定数項・1/s項・s項へ分ける','② P / PI / PD / PIDを判別する','③ Kp, Ki, Kd ↔ Ti, Tdを変換する','④ E(s)=R(s)/{1+C(s)G(s)}','⑤ ステップ1/s・ランプ1/s²を区別','⑥ 定常偏差は lim[s→0] sE(s)','⑦ 時間応答なら閉ループ分母を整理','⑧ 二次形ならTopic 18のωn・ζへ接続']):txt(s,t,.95,2.10+i*.49,4.4,.27,10.5)
txt(s,'安定判別・ボード線図はTopic 20。再粘着はTopic 21。',.96,6.08,4.3,.25,9.2,orange,1)
box(s,5.95,1.42,6.75,5.08);txt(s,'固定過去問ゲート',6.24,1.65,2.4,.28,14,ink,1)
for i,(q,n,k) in enumerate([('H25 一次 機械 問6(2)','1','PID標準形・積分時間'),('R07 二次 機械・制御 問4(1)–(3)','3','PI判別・E(s)・定常偏差'),('H30 二次 機械・制御 問4(1)','2','比例ゲイン・積分時間'),('H23 二次 機械・制御 問4(4)','4','PID名称・Kp/Ti/Td'),('R02 二次 機械・制御 問4(1),(2)','2','P制御・定常偏差・二次形')]):
 y=2.18+i*.65;txt(s,q,6.25,y,2.55,.46,9.2);txt(s,n,8.90,y,.45,.30,11,blue,1,PP_ALIGN.CENTER);txt(s,k,9.48,y,2.85,.42,9.2)
txt(s,'合計5問・12答案要素 → 12 / 12 COVERED',6.28,5.62,5.9,.32,12,green,1,PP_ALIGN.CENTER);txt(s,'固定過去問の個別正答・最終数値は保存しない。',6.30,6.03,5.8,.25,9,muted,0,PP_ALIGN.CENTER)
r.save(O)