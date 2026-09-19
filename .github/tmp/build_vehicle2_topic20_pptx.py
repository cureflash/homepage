from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pathlib import Path
import math, hashlib, zipfile, os
OUT=Path('denken-shinkansen/05_shinkansen_vehicle_2/topics/20_n700_feedback_control')
OUT.mkdir(parents=True, exist_ok=True)
FONT='Noto Sans CJK JP'; W,H=13.333,7.5
ACCENT=RGBColor(27,94,170); ACCENT2=RGBColor(0,128,96); DARK=RGBColor(31,41,55); MID=RGBColor(75,85,99); LIGHT=RGBColor(243,246,249); GRID=RGBColor(214,220,228); RED=RGBColor(180,50,50); ORANGE=RGBColor(196,117,31); WHITE=RGBColor(255,255,255)
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H); blank=prs.slide_layouts[6]

def set_bg(s): s.background.fill.solid(); s.background.fill.fore_color.rgb=WHITE

def text(s,t,x,y,w,h,sz=18,color=DARK,bold=False,align=PP_ALIGN.LEFT):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=bold; r.font.color.rgb=color; return b

def title(s,n,t,sub=''):
    text(s,f'{n:02d}',.45,.28,.7,.4,18,ACCENT,True); text(s,t,1.05,.18,11.6,.5,28,DARK,True)
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.45),Inches(.88),Inches(12.4),Inches(.03)); sh.fill.solid(); sh.fill.fore_color.rgb=ACCENT; sh.line.fill.background()
    if sub:text(s,sub,1.05,.66,11.5,.2,10,MID)

def label(s,t,x,y,w,h=.55,fill=LIGHT,color=DARK,sz=13):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=GRID; tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=True; r.font.color.rgb=color; return sh

def footer(s): text(s,'一般教材モデル。N700系実車の制御定数・伝達関数を示すものではない。',.55,7.16,12.1,.18,9,MID,False,PP_ALIGN.RIGHT)

def arrow(s,x1,y1,x2,y2,color=ACCENT):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=color; ln.line.width=Pt(2); ln.line.end_arrowhead=True; return ln

def chart(s,x,y,w,h,funcs,xmax,ymax,xticks,yticks,legend,marks=None):
    ox=x+.55; oy=y+h-.45; pw=w-.75; ph=h-.75
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(ox),Inches(oy),Inches(ox+pw),Inches(oy)); ln.line.color.rgb=DARK; ln.line.width=Pt(1)
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(ox),Inches(oy),Inches(ox),Inches(oy-ph)); ln.line.color.rgb=DARK; ln.line.width=Pt(1)
    for xv in xticks:
        xx=ox+pw*xv/xmax; ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(xx),Inches(oy),Inches(xx),Inches(oy-ph)); ln.line.color.rgb=GRID; ln.line.width=Pt(.5); text(s,str(xv),xx-.16,oy+.05,.32,.2,8,MID,False,PP_ALIGN.CENTER)
    for yv in yticks:
        yy=oy-ph*yv/ymax; ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(ox),Inches(yy),Inches(ox+pw),Inches(yy)); ln.line.color.rgb=GRID; ln.line.width=Pt(.5); text(s,f'{yv:g}',ox-.4,yy-.1,.35,.2,8,MID,False,PP_ALIGN.RIGHT)
    cols=[ACCENT,ACCENT2,ORANGE,RED]
    N=45
    for j,f in enumerate(funcs):
        pts=[]
        for i in range(N):
            tv=xmax*i/(N-1); val=max(0,min(ymax,f(tv))); pts.append((ox+pw*tv/xmax, oy-ph*val/ymax))
        for a,b in zip(pts,pts[1:]):
            ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1])); ln.line.color.rgb=cols[j]; ln.line.width=Pt(1.6)
        lx=ox+.18+j*1.18; ly=y+.08
        ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(lx),Inches(ly+.08),Inches(lx+.28),Inches(ly+.08)); ln.line.color.rgb=cols[j]; ln.line.width=Pt(2)
        text(s,legend[j],lx+.32,ly-.03,.78,.22,8,MID)
    if marks:
        for tx,ty,lab in marks:
            xx=ox+pw*tx/xmax; yy=oy-ph*ty/ymax
            sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(xx-.045),Inches(yy-.045),Inches(.09),Inches(.09)); sh.fill.solid(); sh.fill.fore_color.rgb=RED; sh.line.fill.background(); text(s,lab,xx+.08,yy-.18,1.25,.4,9,RED,True)
    text(s,'t [s]',ox+pw-.4,oy+.24,.55,.2,8,MID)
    text(s,'y',ox-.42,y+.08,.22,.2,8,MID)

s=prs.slides.add_slide(blank); set_bg(s); title(s,1,'フィードバック制御をブロック線図へ落とす','電験二種：伝達関数・開ループ・閉ループ・偏差')
for txt0,x0,w0 in [('目標値 R',.65,1.35),('比較点',2.35,1.2),('制御要素 C(s)',4.05,1.75),('制御対象 P(s)',6.35,1.75),('出力 Y',8.7,1.3)]: label(s,txt0,x0,2.15,w0,.68,WHITE if txt0!='比較点' else LIGHT,DARK,14)
for a,b in [(2.0,2.35),(3.55,4.05),(5.8,6.35),(8.1,8.7)]: arrow(s,a,2.49,b,2.49)
label(s,'検出要素 H(s)',6.35,4.08,1.75,.62,LIGHT,DARK,13); arrow(s,9.35,2.83,9.35,4.39,ACCENT2); arrow(s,9.35,4.39,8.1,4.39,ACCENT2); arrow(s,6.35,4.39,2.95,4.39,ACCENT2); arrow(s,2.95,4.39,2.95,2.83,ACCENT2)
text(s,'−',2.74,2.77,.25,.25,18,RED,True,PP_ALIGN.CENTER); text(s,'偏差 E = R − HY',1.85,3.15,2.4,.35,14,ACCENT2,True)
text(s,'試験で区別する3つ',10.35,1.45,2.4,.35,17,DARK,True); text(s,'開ループ\nL(s)=C(s)P(s)H(s)',10.35,1.95,2.35,.78,16); text(s,'閉ループ\nY/R=CP/(1+CPH)',10.35,2.95,2.35,.78,16); text(s,'単位負帰還の偏差\nE/R=1/(1+G)',10.35,3.95,2.35,.78,16)
for t0,x0,w0 in [('直列＝積',.75,1.55),('並列＝和/差',2.55,1.75),('負帰還：1+L(s)=0',4.55,2.35),('入力・出力を先に確定',7.15,2.5)]:label(s,t0,x0,5.4,w0,.58,LIGHT,ACCENT,15)
text(s,'N700系は「指令値→偏差→制御→車両/駆動系→検出値」の一般モデルとして接続する。',.75,6.2,9,.52,15); footer(s)
s=prs.slides.add_slide(blank); set_bg(s); title(s,2,'一次遅れ：時定数とステップ応答を結び付ける','G(s)=K/(Ts+1)  →  y(t)=K(1−e^(−t/T))')
chart(s,.55,1.25,7.4,4.8,[lambda t:1-math.exp(-t)],5,1.05,[0,1,2,3,4,5],[0,.5,1],['y=1−e^(−t)'],[(1,1-math.exp(-1),'t=T：63.2%')])
text(s,'一次遅れの要点',8.35,1.35,4.25,.4,19,DARK,True); text(s,'極：s=−1/T\nT>0 なら左半平面 → 安定\n\nt=T のとき\ny(T)=0.632K\n\nT は「最終値へ到達する時間」ではない。',8.35,1.85,4.2,2.45,18)
label(s,'指定可視化① ステップ応答',8.35,4.55,3.8,.55,LIGHT,ACCENT,14); text(s,'式・初期値・時間軸・定常値を明記し、63.2%点まで同じ式から再生成。',8.35,5.22,4.2,.8,14)
text(s,'例：G(s)=2/(3s+1) → y(t)=2(1−e^(−t/3)),  y(3)=1.264',.75,6.35,11.5,.42,15,DARK,True); footer(s)
s=prs.slides.add_slide(blank); set_bg(s); title(s,3,'閉ループ：ゲイン変更と定常偏差を同じ式で読む','単位負帰還 G_K(s)=K/(s+1)')
chart(s,.4,1.25,6.15,3.55,[lambda t,K=K:K/(1+K)*(1-math.exp(-(1+K)*t)) for K in (.5,1,2)],4,.72,[0,1,2,3,4],[0,.35,.7],['K=0.5','K=1','K=2'])
chart(s,6.7,1.25,6.15,3.55,[lambda t,T=T:1-math.exp(-t/T) for T in (.5,1,2)],4,1.05,[0,1,2,3,4],[0,.5,1],['T=0.5','T=1','T=2'])
label(s,'指定可視化② ゲイン変更',.85,4.78,2.55,.52,LIGHT,ACCENT,13); text(s,'T_K(s)=K/(s+1+K)\ny_K=K/(1+K)[1−e^{−(1+K)t}]\nK↑ → このモデルでは定常値↑・実効時定数↓',.75,5.34,5.9,1.05,14)
label(s,'指定可視化③ 時定数変更',7.1,4.78,2.65,.52,LIGHT,ACCENT,13); text(s,'G_T(s)=1/(Ts+1)\nT=0.5,1,2 s：定常ゲインは同じ1\n一次遅れでは T↓ → 立上りが速い',7,5.34,5.4,1.05,14)
text(s,'定常偏差：E/R=1/(1+G) → 入力種類を確認 → 安定性を確認 → e∞=lim sE(s)',.75,6.53,11.7,.35,15,DARK,True); footer(s)
s=prs.slides.add_slide(blank); set_bg(s); title(s,4,'二次遅れ・過渡応答・安定性：極から判定する','固定5問・19答案要素へ接続する最終チェック')
text(s,'二次遅れ',.75,1.35,2.1,.35,20,DARK,True); text(s,'G(s)=Kωₙ²/(s²+2ζωₙs+ωₙ²)\n極：s=−ζωₙ ± ωₙ√(ζ²−1)',.75,1.82,5.35,.85,17)
label(s,'0<ζ<1：減衰振動',.85,2.95,2,.52,LIGHT,ACCENT,13); label(s,'ζ=1：臨界減衰',3,2.95,1.8,.52,LIGHT,ACCENT2,13); label(s,'ζ>1：非振動収束',4.95,2.95,2,.52,LIGHT,ORANGE,13)
text(s,'三次系の安定条件',7.35,1.35,2.9,.35,20,DARK,True); text(s,'a₃s³+a₂s²+a₁s+a₀=0  (a₃>0)\n\na₂>0, a₁>0, a₀>0\na₂a₁ > a₃a₀',7.35,1.82,5.25,1.35,18); text(s,'例：s³+3s²+2s+K=0\n→ 0<K<6\nK=6 は安定境界。漸近安定に含めない。',7.35,3.35,5.2,1,16,RED,True)
text(s,'時間応答・定常値の解法',.75,4.3,3,.35,19,DARK,True); steps=['① 入力を判定：δ(t), 1(t), t, e^(−at)','② R(s)へ変換し Y(s)=G(s)R(s)','③ 部分分数分解 → 逆ラプラス','④ 初期値・最終値・減衰方向で検算','⑤ 最終値の定理は「収束条件を先に確認」']
for i,t0 in enumerate(steps):text(s,t0,.85,4.83+i*.38,6.05,.34,14)
text(s,'過去問ゲート',7.35,4.55,2.2,.35,19,DARK,True); text(s,'H25一次(1)：安定判定\nR7二次(1)〜(5)：偏差・応答\nR6二次(1)〜(4)：直列/並列・時間応答\nR4二次(1)〜(4)：特性方程式・安定条件\nR3二次(1)〜(5)：開閉ループ・二次/三次安定',7.35,5.02,5.1,1.48,14); footer(s)
p=OUT/'20_n700_feedback_control_images.pptx'; prs.save(p); print(p.stat().st_size,hashlib.sha256(p.read_bytes()).hexdigest())
