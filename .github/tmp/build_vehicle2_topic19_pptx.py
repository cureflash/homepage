from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path

OUT=Path('denken-shinkansen/05_shinkansen_vehicle_2/topics/19_mini_shinkansen_dual_voltage_main_circuit/19_mini_shinkansen_dual_voltage_main_circuit_images.pptx')
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
FONT='Noto Sans CJK JP'
NAVY=RGBColor(25,39,56); BLUE=RGBColor(39,92,154); LIGHT=RGBColor(240,244,248); MID=RGBColor(215,224,234); TEXT=RGBColor(35,42,50); GRAY=RGBColor(92,102,112); WHITE=RGBColor(255,255,255); ACC=RGBColor(220,234,248)

def add_text(slide,x,y,w,h,text,size=18,bold=False,color=TEXT,align=PP_ALIGN.LEFT,fill=None,line=None,margin=.08):
    shp=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is not None:
        shp.fill.solid(); shp.fill.fore_color.rgb=fill
    else: shp.fill.background()
    shp.line.fill.background() if line is None else None
    if line is not None:
        shp.line.color.rgb=line; shp.line.width=Pt(1)
    tf=shp.text_frame; tf.clear(); tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return shp

def box(slide,x,y,w,h,text,size=17,fill=LIGHT,bold=True):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill; shp.line.color.rgb=MID; shp.line.width=Pt(1)
    tf=shp.text_frame; tf.clear(); tf.margin_left=tf.margin_right=Inches(.08); tf.margin_top=tf.margin_bottom=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=text; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=TEXT
    return shp

def line(slide,x1,y1,x2,y2,width=2,color=BLUE):
    l=slide.shapes.add_connector(1, Inches(x1),Inches(y1),Inches(x2),Inches(y2)); l.line.color.rgb=color; l.line.width=Pt(width); return l

def arrow(slide,x1,y1,x2,y2):
    l=line(slide,x1,y1,x2,y2,2,BLUE); l.line.end_arrowhead=True; return l

def title(slide,txt,sub=None):
    add_text(slide,.45,.18,12.45,.48,txt,24,True,NAVY)
    if sub: add_text(slide,.48,.67,12.2,.32,sub,11,False,GRAY)
    line(slide,.45,1.02,12.88,1.02,1,MID)

def footer(slide,n):
    add_text(slide,.48,7.08,10.6,.18,'教材用一般モデル（未確認の実車仕様は真値化しない）',9,False,GRAY)
    add_text(slide,12.15,7.05,.55,.2,f'{n}/4',9,False,GRAY,PP_ALIGN.RIGHT)

s=prs.slides.add_slide(blank); title(s,'Topic 19 複電圧主回路：試験論点を一つの流れで見る','第二種一次（機械・電力）＋二次（機械・制御）／固定5問・9答案要素')
box(s,.55,1.35,1.75,.85,'供給条件\n複電圧\nV1 が変わる',15,ACC)
box(s,2.6,1.35,2.2,.85,'主変圧器\na=N1/N2=V1/V2\n=I2/I1',14)
box(s,5.1,1.35,2.0,.85,'主変換装置\nVo = D Vd',15)
box(s,7.4,1.35,1.8,.85,'主負荷\n必要な電圧・出力',15)
for a,b in [((2.3,1.78),(2.6,1.78)),((4.8,1.78),(5.1,1.78)),((7.1,1.78),(7.4,1.78))]: arrow(s,*a,*b)
box(s,5.1,2.55,2.0,.7,'補助電源\n補機へ必要条件',14); arrow(s,6.1,2.2,6.1,2.55)
box(s,.65,3.55,5.3,.8,'電源切替（機能ブロック）\n供給条件の識別 → 適合する変圧・変換条件の選択',15,ACC)
box(s,6.25,3.55,5.95,.8,'絶縁協調\n想定過電圧・保護レベル・機器絶縁強度を系統として整合',15,ACC)
add_text(s,.72,4.55,5.2,.38,'※実車の接点構成・切替順序は断定しない',11,False,GRAY)
add_text(s,.55,5.05,12.0,.38,'固定答案要素の接続',17,True,NAVY)
labels=[('巻数比・電圧比','R5一次 機械 問3 (2)'),('通流率・出力調整','R5一次 機械 問4 (4)(5)'),('絶縁協調','R5一次 電力 問6 (3)'),('タップ切換一般原理','H25一次 電力 問2 (1)(2)(3)'),('電圧・電流換算','R5二次 機械・制御 問2 (1)(2)')]
for i,(a,b) in enumerate(labels):
    x=.55+i*2.48; box(s,x,5.55,2.25,.48,a,11,WHITE,True); add_text(s,x,6.08,2.25,.42,b,9,False,GRAY,PP_ALIGN.CENTER)
footer(s,1)

s=prs.slides.add_slide(blank); title(s,'電源電圧―一次電流：同一出力では I1 ∝ 1/V1','条件：Pout=6.84 MW、η=0.95、cosφ=0.90、単相等価（すべて教材用仮定値）')
add_text(s,.55,1.35,3.2,.5,'I1 = Pout /(η V1 cosφ)',20,True,NAVY,PP_ALIGN.CENTER,ACC,MID)
add_text(s,.7,2.0,2.6,.85,'条件A → 条件B\n10 kV → 800 A\n8 kV → 1000 A',16,True,TEXT,PP_ALIGN.CENTER,LIGHT,MID)
add_text(s,.7,3.0,2.6,.55,'I1,B / I1,A = 10/8 = 1.25',15,True,BLUE,PP_ALIGN.CENTER)
add_text(s,.55,4.05,3.0,1.0,'試験での注意\n反比例を使えるのは Pout・η・cosφ が同一のとき。低電圧条件ほど必要電流が増える。',12,False,TEXT,PP_ALIGN.LEFT,LIGHT,MID)
cx,cy,cw,ch=4.1,1.55,8.1,4.95
line(s,cx,cy+ch,cx+cw,cy+ch,1,NAVY); line(s,cx,cy,cx,cy+ch,1,NAVY)
vals=[(8,1000),(10,800),(12.5,640),(16,500),(20,400)]
for yv in [400,600,800,1000]:
    yy=cy+ch-(yv-300)/(1050-300)*ch; line(s,cx,yy,cx+cw,yy,.6,MID); add_text(s,cx-.62,yy-.12,.52,.24,str(yv),9,False,GRAY,PP_ALIGN.RIGHT)
for xv in [8,10,12.5,16,20]:
    xx=cx+(xv-8)/12*cw; add_text(s,xx-.35,cy+ch+.08,.7,.25,str(xv),9,False,GRAY,PP_ALIGN.CENTER)
pts=[]
for xv,yv in vals:
    xx=cx+(xv-8)/12*cw; yy=cy+ch-(yv-300)/750*ch; pts.append((xx,yy))
for (x1,y1),(x2,y2) in zip(pts,pts[1:]): line(s,x1,y1,x2,y2,2.5,BLUE)
for xx,yy in pts:
    dot=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xx-.07),Inches(yy-.07),Inches(.14),Inches(.14)); dot.fill.solid(); dot.fill.fore_color.rgb=BLUE; dot.line.color.rgb=BLUE
add_text(s,7.3,6.65,2.0,.25,'V1 [kV]',10,True,GRAY,PP_ALIGN.CENTER); add_text(s,3.62,1.15,.65,.25,'I1 [A]',10,True,GRAY); footer(s,2)

s=prs.slides.add_slide(blank); title(s,'タップ条件―二次電圧：実効巻数を式へ入れる','条件：V1=10 kV、N2=200 turn、理想変圧器（教材用一般モデル）')
add_text(s,.55,1.35,3.15,.52,'V2 = V1 N2 / N1,tap',20,True,NAVY,PP_ALIGN.CENTER,ACC,MID)
add_text(s,.65,2.1,2.85,.85,'向きの確認\nN1,tap を減らす → V2 は上がる\nN1,tap を増やす → V2 は下がる',14,True,TEXT,PP_ALIGN.CENTER,LIGHT,MID)
add_text(s,.55,3.35,3.2,1.25,'H25一次「電力」問2への接続\n隣接タップを橋絡すると電位差により循環電流が生じ得る。限流インピーダンスで抑える。\n※一般原理であり、実車採用方式とは断定しない。',11,False,TEXT,PP_ALIGN.LEFT,LIGHT,MID)
cx,cy,cw,ch=4.35,1.55,7.6,4.85; line(s,cx,cy+ch,cx+cw,cy+ch,1,NAVY); line(s,cx,cy,cx,cy+ch,1,NAVY)
data=[('1250',1.6),('1000',2.0),('800',2.5)]
for yv in [0,1,2,3]:
    yy=cy+ch-yv/3*ch; line(s,cx,yy,cx+cw,yy,.6,MID); add_text(s,cx-.45,yy-.12,.35,.24,str(yv),9,False,GRAY,PP_ALIGN.RIGHT)
for i,(lab,v) in enumerate(data):
    bx=cx+.75+i*2.15; bh=v/3*ch; rect=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(cy+ch-bh), Inches(1.1), Inches(bh)); rect.fill.solid(); rect.fill.fore_color.rgb=BLUE; rect.line.color.rgb=BLUE
    add_text(s,bx,cy+ch-bh-.35,1.1,.3,f'{v:.1f}',11,True,BLUE,PP_ALIGN.CENTER); add_text(s,bx-.1,cy+ch+.08,1.3,.28,lab,10,False,GRAY,PP_ALIGN.CENTER)
add_text(s,7.0,6.65,2.5,.25,'N1,tap [turn]',10,True,GRAY,PP_ALIGN.CENTER); add_text(s,3.9,1.15,.65,.25,'V2 [kV]',10,True,GRAY); footer(s,3)

s=prs.slides.add_slide(blank); title(s,'同一出力時の電源条件比較：変圧・変換・絶縁を分けて判断する','固定式系だけで比較。実車定格・実車切替方式はここから決めない。')
add_text(s,.55,1.35,5.65,.38,'条件A / 条件B',18,True,NAVY)
headers=['項目','条件A','条件B']; xs=[.55,2.55,4.15]; ws=[2.0,1.6,1.6]
for x,w,h in zip(xs,ws,headers): box(s,x,1.8,w,.45,h,12,ACC,True)
rows=[('V1 [kV]','10','8'),('Pout [MW]','6.84','6.84'),('η','0.95','0.95'),('cosφ','0.90','0.90'),('I1 [A]','800','1000'),('I1/I1,A','1.00','1.25')]
for ri,row in enumerate(rows):
    yy=2.28+ri*.52
    for ci,val in enumerate(row): add_text(s,xs[ci],yy,ws[ci],.45,val,12,ci==0,TEXT,PP_ALIGN.CENTER,WHITE,MID)
add_text(s,6.55,1.35,5.9,.38,'主変換装置：目標出力電圧へ調整',18,True,NAVY)
box(s,6.75,1.95,2.2,.7,'条件A\nVd=1200 V',14,LIGHT,True); box(s,10.0,1.95,2.2,.7,'条件B\nVd=900 V',14,LIGHT,True)
add_text(s,7.0,2.85,1.7,.45,'D=720/1200\n=0.60',15,True,BLUE,PP_ALIGN.CENTER); add_text(s,10.25,2.85,1.7,.45,'D=720/900\n=0.80',15,True,BLUE,PP_ALIGN.CENTER)
add_text(s,8.75,3.55,1.45,.38,'Vo = D Vd',18,True,NAVY,PP_ALIGN.CENTER,ACC,MID)
add_text(s,.55,5.62,3.75,.78,'変圧・変換\n巻数比・タップ・通流率は所要電圧条件へ合わせる。',12,False,TEXT,PP_ALIGN.CENTER,LIGHT,MID)
add_text(s,4.8,5.62,3.75,.78,'絶縁協調\n想定過電圧と機器絶縁強度を系統として整合する。',12,False,TEXT,PP_ALIGN.CENTER,LIGHT,MID)
add_text(s,9.05,5.62,3.75,.78,'電源切替\n適合する変圧・変換条件を選ぶ機能として扱う。',12,False,TEXT,PP_ALIGN.CENTER,LIGHT,MID)
add_text(s,.62,6.55,12.0,.32,'3つを一つの未確認実車仕様へ混同しない。',13,True,BLUE,PP_ALIGN.CENTER); footer(s,4)

OUT.parent.mkdir(parents=True, exist_ok=True); prs.save(OUT)
print(OUT, OUT.stat().st_size)
