from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pathlib import Path

base=Path('denken-shinkansen/05_shinkansen_vehicle_2/topics/34_compressor_pump_fan_load_characteristics')
out=base/'34_compressor_pump_fan_load_characteristics_images.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
NAVY=RGBColor(20,35,58); BLUE=RGBColor(42,98,171); CYAN=RGBColor(63,157,190); TEAL=RGBColor(35,128,123); GREEN=RGBColor(67,145,91); ORANGE=RGBColor(221,132,54); RED=RGBColor(190,72,64); PURPLE=RGBColor(113,82,161); INK=RGBColor(31,41,55); GRAY=RGBColor(92,105,120); LIGHT=RGBColor(243,246,249); MID=RGBColor(220,226,233); WHITE=RGBColor(255,255,255); PALE_BLUE=RGBColor(231,240,250); PALE_CYAN=RGBColor(230,246,249); PALE_GREEN=RGBColor(234,246,238); PALE_ORANGE=RGBColor(252,241,229); PALE_RED=RGBColor(250,235,233)
FONT='Noto Sans CJK JP'; MONO='Noto Sans Mono CJK JP'

def bg(s,c=WHITE):
    f=s.background.fill; f.solid(); f.fore_color.rgb=c

def rect(s,x,y,w,h,fill=WHITE,line=MID,r=True,lw=1.2):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line; q.line.width=Pt(lw); return q

def ln(s,x1,y1,x2,y2,c=GRAY,lw=1.5,arrow=False):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(lw)
    if arrow:
        e=q._element.spPr.ln; t=OxmlElement('a:tailEnd'); t.set('type','none'); h=OxmlElement('a:headEnd'); h.set('type','triangle'); e.append(t); e.append(h)
    return q

def text(s,x,y,w,h,v,size=18,c=INK,b=False,align=PP_ALIGN.LEFT,font=FONT,valign=MSO_ANCHOR.MIDDLE,margin=.06):
    q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=q.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=v; r.font.name=font; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=c; return q

def bullets(s,x,y,w,h,items,size=14,fill=WHITE,linec=MID,title=None,tc=NAVY):
    q=rect(s,x,y,w,h,fill,linec,True,1); tf=q.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.16); tf.margin_right=Inches(.12); tf.margin_top=Inches(.11); tf.margin_bottom=Inches(.08)
    p=tf.paragraphs[0]; p.text=title or ''; p.font.name=FONT; p.font.size=Pt(size+1); p.font.bold=True; p.font.color.rgb=tc; p.space_after=Pt(4)
    for it in items:
        p=tf.add_paragraph(); p.text='• '+it; p.font.name=FONT; p.font.size=Pt(size); p.font.color.rgb=INK; p.space_after=Pt(2)
    return q

def titlebar(s,n,t,sub=''):
    badge=rect(s,.55,.25,.55,.42,BLUE,BLUE,True,0); text(s,.55,.25,.55,.42,f'{n:02d}',16,WHITE,True,PP_ALIGN.CENTER)
    text(s,1.2,.18,11.5,.55,t,24,NAVY,True); text(s,1.22,.70,11.2,.28,sub,10,GRAY) if sub else None; ln(s,.55,1.05,12.75,1.05,MID,1)

def pill(s,x,y,w,v,c,size=11): rect(s,x,y,w,.38,c,c,True,0); text(s,x,y,w,.38,v,size,WHITE,True,PP_ALIGN.CENTER)
def formula(s,x,y,w,h,v,fill=LIGHT,c=NAVY,size=20): rect(s,x,y,w,h,fill,fill,True,0); text(s,x+.08,y+.03,w-.16,h-.06,v,size,c,True,PP_ALIGN.CENTER,MONO)

def notes(s,urls):
    try: s.notes_slide.notes_text_frame.text='Sources:\n'+'\n'.join(urls)
    except Exception: pass

OFFICIAL='https://www.shiken.or.jp/chief/second/qa/'; R7Q='https://www.shiken.or.jp/chief/upload/20251116_ch_second_q02.pdf'; R7A='https://www.shiken.or.jp/chief/upload/20251116_ch_second_a01_tk.pdf'; R4Q='https://www.shiken.or.jp/chief/upload/20221113_ch_second_q02.pdf'; R4A='https://www.shiken.or.jp/chief/upload/20221113_ch_second_a01.pdf'; H25Q='https://www.shiken.or.jp/chief/upload/20130831_ch_second_q03.pdf'; H25A='https://www.shiken.or.jp/chief/upload/20130831_ch_second_a01.pdf'; ESYSM='https://e-sysnet.com/%E9%9B%BB%E5%8B%95%E6%A9%9F%E3%81%AE%E5%88%B6%E5%8B%95%E3%81%A8%E6%89%80%E8%A6%81%E5%87%BA%E5%8A%9B/'; ESYSP='https://e-sysnet.com/%E3%83%9D%E3%83%B3%E3%83%97%E7%94%A8%E3%83%BB%E9%80%81%E9%A2%A8%E6%A9%9F%E3%81%AE%E9%81%8B%E8%BB%A2%E3%81%A8%E5%81%9C%E6%AD%A2/'

s=prs.slides.add_slide(blank); bg(s); titlebar(s,1,'補機負荷を「電力 → トルク → 流体仕事」でつなぐ','Topic 34｜コンプレッサ・ポンプ・ファンの負荷特性｜固定5問・27答案要素')
nodes=[('SIV / 補助電源',.7,PALE_BLUE,BLUE),('インバータ',3.25,PALE_CYAN,TEAL),('電動機',5.65,PALE_GREEN,GREEN),('補機負荷',8.05,PALE_ORANGE,ORANGE),('空気・液体\nの仕事',10.55,PALE_RED,RED)]
for i,(lab,x,fc,lc) in enumerate(nodes):
    rect(s,x,2,1.95,1.05,fc,lc,True,1.6); text(s,x+.08,2.1,1.79,.85,lab,15,lc,True,PP_ALIGN.CENTER)
    if i<4: ln(s,x+1.95,2.52,nodes[i+1][1],2.52,GRAY,2,True)
formula(s,.85,3.45,2.55,.72,'P = Tω',PALE_BLUE,BLUE,22); formula(s,3.62,3.45,3.3,.72,'η = Pout / Pin',PALE_GREEN,GREEN,19); formula(s,7.15,3.45,4.85,.72,'Ns = 120f / p   ,   s = (Ns−N)/Ns',PALE_CYAN,TEAL,17)
text(s,.75,4.55,2.2,.35,'SPEC固定 7項目',13,NAVY,True)
x=.75
for lab,c in [('電動機出力',BLUE),('負荷トルク',PURPLE),('ポンプ動力',TEAL),('ファン動力',ORANGE),('回転速度',GREEN),('効率',RED),('インバータ制御',NAVY)]:
    w=max(1.15,len(lab)*.22+.55); pill(s,x,5,w,lab,c,10); x+=w+.12
bullets(s,.75,5.7,11.95,1.05,['新幹線実車の未確認定格値は使わず、電験二種の一般モデルで解析する。','コンプレッサは軸出力・トルク・速度・効率・インバータ制御まで。圧縮熱力学・冷凍サイクルは追加しない。'],12,LIGHT,MID)
notes(s,[OFFICIAL,H25Q,H25A,ESYSM,ESYSP])

s=prs.slides.add_slide(blank); bg(s); titlebar(s,2,'電動機出力・負荷トルク・動作点','P=Tω / 同期速度 / すべり / 停動トルクを同じ速度軸で読む')
formula(s,.7,1.25,3.4,.62,'P = Tω = 2πNT / 60',PALE_BLUE,BLUE,18); formula(s,4.35,1.25,3.1,.62,'J dω/dt = Tm − TL',PALE_GREEN,GREEN,17); formula(s,7.7,1.25,4.2,.62,'Ns=120f/p  ,  s=(Ns−N)/Ns',PALE_CYAN,TEAL,16)
x0,y0,w,h=.85,5.85,6.35,3.45; ln(s,x0,y0,x0+w,y0,NAVY,1.5,True); ln(s,x0,y0,x0,y0-h,NAVY,1.5,True); text(s,x0+w-1.1,y0+.12,1.15,.3,'回転速度 N',11,NAVY,True,PP_ALIGN.RIGHT); text(s,x0-.62,y0-h-.15,.55,.3,'T',12,NAVY,True,PP_ALIGN.CENTER)
pts=[]
for i in range(31):
    u=i/30; yy=.22+2.55*(u**.7)*(1-u)**.55; pts.append((x0+.15+u*(w-.4),y0-.25-yy))
for a,b in zip(pts,pts[1:]): ln(s,a[0],a[1],b[0],b[1],BLUE,2.4)
pts=[]
for i in range(25):
    u=i/24; pts.append((x0+.15+u*(w-.4),y0-.22-2*u*u))
for a,b in zip(pts,pts[1:]): ln(s,a[0],a[1],b[0],b[1],ORANGE,2.4)
px=x0+.15+.73*(w-.4); py=y0-.22-2*(.73**2); q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(px-.07),Inches(py-.07),Inches(.14),Inches(.14)); q.fill.solid(); q.fill.fore_color.rgb=RED; q.line.fill.background(); text(s,px+.1,py-.18,1.25,.38,'定常動作点\nTm = TL',10,RED,True); text(s,1.25,2.18,1.55,.35,'電動機トルク',11,BLUE,True); text(s,2.75,4.95,1.4,.35,'ファン負荷 ∝N²',11,ORANGE,True)
bullets(s,7.62,2.12,4.65,3.75,['分子 3V₁²(r₂′/s) を含むトルク式','電動機運転：s > 0','最大トルク＝停動トルク','理想化式では最大トルク値は r₂′ に依存しない','同期速度近傍：T ∝ s','定常：Tm=TL／過負荷耐量は問題条件で判定'],12,WHITE,BLUE,'R7一次 問2 ＋ H25一次 接続',BLUE)
notes(s,[OFFICIAL,H25Q,H25A,'https://denken-ou.com/c2/kikair7-2/'])

s=prs.slides.add_slide(blank); bg(s); titlebar(s,3,'効率と損失を、入力から軸出力まで分解する','R7二次「機械・制御」問2の固定6答案要素を数値で再現')
flow=[('入力\n25.14 kW',.75,PALE_BLUE,BLUE),('一次銅損\n0.798 kW',3.05,PALE_RED,RED),('固定損\n1.55 kW',5.35,PALE_ORANGE,ORANGE),('二次銅損\n0.798 kW',7.65,PALE_RED,RED),('軸出力\n22.0 kW',9.95,PALE_GREEN,GREEN)]
for i,(lab,x,fc,lc) in enumerate(flow):
    rect(s,x,2.05,1.75,1,fc,lc,True,1.4); text(s,x+.05,2.12,1.65,.86,lab,13,lc,True,PP_ALIGN.CENTER)
    if i<4: ln(s,x+1.75,2.55,flow[i+1][1],2.55,GRAY,1.7,True)
for x,y,w,v,c in [(.78,3.45,3.15,'Ns = 120×60/6 = 1200 min⁻¹',TEAL),(.78,4.18,3.15,'s₁ = (1200−1158)/1200 = 0.035',TEAL),(.78,4.91,3.15,'T₁ = 22 kW / ω = 181 N·m',BLUE),(4.18,3.45,3.45,'Pc₂ = s/(1−s) × 22 kW = 798 W',RED),(4.18,4.18,3.45,'Pin = 22/0.875 = 25.14 kW',GREEN),(4.18,4.91,3.45,'Pf ≈ 3.143 − 2×0.798 = 1.55 kW',ORANGE)]: formula(s,x,y,w,.58,v,LIGHT,c,13 if x>1 else 14)
rect(s,7.9,3.4,4.4,2.15,LIGHT,MID,True,1); text(s,8.1,3.6,3.95,.35,'半トルク時（同期速度近傍 T∝s）',13,NAVY,True); text(s,8.1,4.1,3.9,.32,'s₂ = 0.0175  →  N₂ = 1179 ≈ 1180 min⁻¹',12); text(s,8.1,4.54,3.9,.32,'P₂ ≈ 0.5T₁ × 2πN₂/60 ≈ 11.2 kW',12); text(s,.8,6.25,11.6,.62,'検算：η = Pout/Pin。効率を掛ける向き、軸出力と内部機械出力、min⁻¹→rad/s の変換を混同しない。',12,GRAY,True,PP_ALIGN.CENTER)
notes(s,[R7Q,R7A,'https://denken-ou.com/c2/kikaiseigyor7-2/'])

s=prs.slides.add_slide(blank); bg(s); titlebar(s,4,'ポンプ・ファン負荷：速度を下げると必要動力は大きく下がる','三種既習式は「負荷側→電動機軸出力」の接続に限定。二種固定過去問の直接出題式とは扱わない。')
rect(s,.65,1.35,5.45,2.15,PALE_CYAN,TEAL,True,1.1); text(s,.9,1.53,2.4,.38,'ポンプ：仕事率から逆算',15,TEAL,True); formula(s,.9,2.04,2.2,.55,'Ph = ρgQH',WHITE,TEAL,17); formula(s,3.25,2.04,2.45,.55,'Pshaft = Ph / ηp',WHITE,TEAL,15); text(s,.95,2.78,4.85,.43,'さらに Pin = Pshaft / ηm。上流へ戻るほど損失分だけ必要入力は増える。',11)
rect(s,6.4,1.35,6.25,2.15,PALE_ORANGE,ORANGE,True,1.1); text(s,6.65,1.53,3.1,.38,'ファン相似則（成立条件つき）',15,ORANGE,True); formula(s,6.7,2.04,5.3,.55,'Q∝N   Δp∝N²   T∝N²   P∝N³',WHITE,ORANGE,16); text(s,6.72,2.78,5.25,.43,'同一ファン・同一流体・幾何形状不変・相似運転の範囲で使う。',11)
text(s,.75,3.83,3.1,.36,'例：60 Hz → 45 Hz（速度比 k=0.75）',14,NAVY,True); basey=6.35; ln(s,1,basey,7.2,basey,NAVY,1.2)
for x,label,val,c in [(1.35,'速度',1,BLUE),(2.15,'速度',.75,CYAN),(3.45,'軸動力',1,ORANGE),(4.25,'軸動力',.421875,RED)]:
    hh=1.65*val; rect(s,x,basey-hh,.58,hh,c,c,False,0); text(s,x-.1,basey-hh-.32,.78,.28,f'{val*100:.1f}%',10,c,True,PP_ALIGN.CENTER); text(s,x-.22,basey+.05,1.02,.45,label,10,INK,False,PP_ALIGN.CENTER)
text(s,1.15,6.93,1.9,.25,'60 Hz    45 Hz',9,GRAY,False,PP_ALIGN.CENTER); text(s,3.25,6.93,1.9,.25,'18.0 kW  7.59 kW',9,GRAY,False,PP_ALIGN.CENTER)
rect(s,7.5,3.8,5.05,2.65,LIGHT,MID,True,1); text(s,7.78,4.02,4.5,.35,'VVVF速度制御が省エネにつながる理由',14,NAVY,True); formula(s,7.78,4.54,4.45,.55,'k=0.75 → P₂/P₁ = k³ = 0.4219',PALE_ORANGE,ORANGE,14); text(s,7.8,5.23,4.38,.55,'18 kW → 7.59 kW（仮定例）\nηm=88% なら入力 8.63 kW',12,INK,True,PP_ALIGN.CENTER); pill(s,8.22,5.95,3.55,'V/f一定：45 Hz時は電圧も0.75倍',TEAL,10)
notes(s,[H25Q,H25A,ESYSP])

s=prs.slides.add_slide(blank); bg(s); titlebar(s,5,'インバータ制御：周波数・電圧・速度・トルクを一つの流れで追う','R6一次・H25一次・R4二次の答案要素を回路ブロックへ配置')
blocks=[('三相交流\n220 V',.55,PALE_BLUE,BLUE),('整流器',2.65,PALE_CYAN,TEAL),('DC link\n297 V',4.55,LIGHT,NAVY),('PWM\nインバータ',6.65,PALE_ORANGE,ORANGE),('誘導電動機',8.85,PALE_GREEN,GREEN),('補機負荷',11.05,PALE_RED,RED)]
for i,(lab,x,fc,lc) in enumerate(blocks):
    rect(s,x,1.45,1.55,.95,fc,lc,True,1.3); text(s,x+.04,1.51,1.47,.82,lab,12,lc,True,PP_ALIGN.CENTER)
    if i<5: ln(s,x+1.55,1.92,blocks[i+1][1],1.92,GRAY,1.5,True)
bullets(s,.6,2.83,4.08,2.85,['Ed = 1.35×220 = 297 V','Id = 10 kW / 297 = 33.7 A','I₁ = 10 kW /(√3×220) = 26.2 A','Vph,peak = (Ed/2)×0.9 ≈ 134 V','VLL,rms ≈ 164 V','fを1/2 → 電圧も1/2（V/f一定）'],11,WHITE,TEAL,'R4二次 問3：固定6要素',TEAL)
bullets(s,4.88,2.83,3.45,2.85,['誘導電動機','回生：N>Ns → s<0','ブラシ不要','主変換素子：IGBT','低速〜基底速度：V/f一定'],12,WHITE,BLUE,'R6一次 問2：固定5要素',BLUE)
bullets(s,8.53,2.83,4.1,2.85,['P=Tω：ωは回転角速度','過負荷耐量','定常：Tm=TL','ベクトル制御：トルク成分','θsl = ∫ωsl dt'],12,WHITE,PURPLE,'H25一次 問3：固定5要素',PURPLE)
formula(s,1.55,6.15,10.15,.58,'周波数 f ↓ → 同期速度 Ns=120f/p ↓ → 回転速度を制御 → 負荷側の必要動力へ接続',LIGHT,NAVY,14); notes(s,[OFFICIAL,R4Q,R4A,H25Q,H25A,'https://denken-ou.com/c2/kikair6-2/'])

s=prs.slides.add_slide(blank); bg(s); titlebar(s,6,'試験での解法順序と品質ゲート','固定5問・27答案要素 / SPEC 7項目をPowerPoint内で全接続')
steps=[('1','負荷側','必要軸出力 / 負荷トルク\nポンプ・ファンは成立条件を確認',TEAL),('2','電動機','P=Tω、Ns、s、動作点\n効率・損失を分解',BLUE),('3','変換器','整流 → DC link → PWM\nV/f と周波数を追う',ORANGE),('4','検算','単位・符号・入力/出力\n公式条件と整合',GREEN)]
for i,(n,hd,body,c) in enumerate(steps):
    x=.65+i*3.12; pill(s,x,1.38,.55,n,c,12); rect(s,x,1.9,2.82,1.1,WHITE,c,True,1.2); text(s,x+.12,2.03,2.55,.28,hd,13,c,True,PP_ALIGN.CENTER); text(s,x+.12,2.38,2.55,.47,body,10,INK,False,PP_ALIGN.CENTER); ln(s,x+2.82,2.45,x+3.1,2.45,GRAY,1.3,True) if i<3 else None
text(s,.72,3.38,4,.36,'固定過去問 → PowerPoint',14,NAVY,True)
for i,(q,cov,sl) in enumerate([('R7一次 機械 問2','5 / 5','S2'),('R6一次 機械 問2','5 / 5','S5'),('H25一次 機械 問3','5 / 5','S2・S4・S5'),('R7二次 機械・制御 問2','6 / 6','S3'),('R4二次 機械・制御 問3','6 / 6','S5')]):
    y=3.82+i*.48; rect(s,.72,y,5.9,.42,LIGHT if i%2==0 else WHITE,MID,False,.5); text(s,.84,y+.02,3.45,.36,q,10); text(s,4.32,y+.02,.95,.36,cov,10,GREEN,True,PP_ALIGN.CENTER); text(s,5.3,y+.02,1.15,.36,sl,10,BLUE,True,PP_ALIGN.CENTER)
text(s,7.05,3.38,2.4,.36,'SPEC固定 7 / 7',14,NAVY,True)
for i,(a,b) in enumerate([('電動機出力','S2・S3'),('負荷トルク','S2'),('ポンプ動力','S4'),('ファン動力','S4'),('回転速度','S2・S3・S5'),('効率','S3・S4'),('インバータ制御','S4・S5')]):
    col=i%2; row=i//2; x=7.05+col*2.75; y=3.82+row*.62; rect(s,x,y,2.55,.5,PALE_GREEN,GREEN,True,.8); text(s,x+.12,y+.03,1.55,.42,a,10,GREEN,True); text(s,x+1.62,y+.03,.76,.42,b,9,GRAY,True,PP_ALIGN.RIGHT)
bullets(s,7.05,6.28,5.55,.8,['未確認実車値なし／圧縮熱力学・冷凍サイクルなし／Topic 35以降の先取りなし'],10,PALE_RED,RED,'境界',RED); notes(s,[OFFICIAL,R7Q,R7A,R4Q,R4A,H25Q,H25A])

prs.core_properties.title='Topic 34 コンプレッサ・ポンプ・ファンの負荷特性'; prs.core_properties.subject='電験二種 新幹線車両教材 解説画像PowerPoint'; prs.core_properties.author='OpenAI'; prs.save(out); print(out)
