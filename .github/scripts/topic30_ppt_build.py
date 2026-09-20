from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
import matplotlib.pyplot as plt
import numpy as np
import os

TOPIC='denken-shinkansen/05_shinkansen_vehicle_2/topics/30_l0_inductive_power_collection_contactless_power_transfer'
OUT=os.environ.get('TOPIC30_OUT', TOPIC)
os.makedirs(OUT, exist_ok=True)
plt.rcParams['font.family']='Noto Sans CJK JP'
plt.rcParams['axes.unicode_minus']=False

def chart(path,x,y,title,xlabel,ylabel,logx=False,vline=None):
    fig,ax=plt.subplots(figsize=(4.1,2.65),dpi=180)
    ax.plot(x,y,marker='o',lw=2.2)
    if logx:
        ax.set_xscale('log',base=2); ax.set_xticks(x); ax.get_xaxis().set_major_formatter(plt.ScalarFormatter())
    if vline is not None: ax.axvline(vline,ls='--',lw=1,alpha=.7)
    ax.set_xlabel(xlabel,fontsize=10); ax.set_ylabel(ylabel,fontsize=10); ax.set_title(title,fontsize=12,pad=7)
    ax.grid(True,alpha=.25); ax.tick_params(labelsize=8); fig.tight_layout(); fig.savefig(path,bbox_inches='tight',facecolor='white'); plt.close(fig)

charts=[os.path.join('/tmp',f'topic30_chart{i}.png') for i in range(1,4)]
chart(charts[0],np.array([.10,.20,.30,.40,.50]),np.array([25.62,102.49,230.59,409.94,640.54]),'結合係数―伝送電力','結合係数 k','負荷電力 $P_L$ [W]')
chart(charts[1],np.array([.5,1,2,4,8,16]),np.array([60.68,73.35,80.05,80.10,73.49,60.91]),'負荷条件―効率','負荷抵抗 $R_L$ [Ω]','効率 η [%]',True)
chart(charts[2],np.array([500,750,1000,1250,1500]),np.array([4.88,49.23,230.59,182.63,119.65]),'周波数特性','周波数 f [Hz]','負荷電力 $P_L$ [W]',False,1000)

prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
BG=RGBColor(248,250,252); NAVY=RGBColor(19,42,63); TEAL=RGBColor(0,126,140); BLUE=RGBColor(61,109,160); ORANGE=RGBColor(211,118,42); GRAY=RGBColor(91,105,120); LIGHT=RGBColor(230,237,243); LIGHT2=RGBColor(239,244,247); WHITE=RGBColor(255,255,255); RED=RGBColor(178,58,58); FONT='Noto Sans CJK JP'

def bg(s):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height); sh.fill.solid(); sh.fill.fore_color.rgb=BG; sh.line.fill.background(); s.shapes._spTree.remove(sh._element); s.shapes._spTree.insert(2,sh._element)

def text(s,t,x,y,w,h,size=16,color=NAVY,bold=False,align=PP_ALIGN.LEFT):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(.05)
    p=tf.paragraphs[0]; p.text=t; p.alignment=align; p.font.name=FONT; p.font.size=Pt(size); p.font.color.rgb=color; p.font.bold=bold; return tb

def title(s,t,sec):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(13.333),Inches(.12)); sh.fill.solid(); sh.fill.fore_color.rgb=TEAL; sh.line.fill.background()
    text(s,t,.58,.28,11.7,.58,25,NAVY,True); text(s,sec,11.65,.35,1.05,.35,10,TEAL,True,PP_ALIGN.RIGHT)

def card(s,lines,x,y,w,h,fill=WHITE,line=LIGHT):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line
    tb=s.shapes.add_textbox(Inches(x+.12),Inches(y+.08),Inches(w-.24),Inches(h-.16)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(.12)
    for i,(t,sz,b,c,a) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=t; p.font.name=FONT; p.font.size=Pt(sz); p.font.bold=b; p.font.color.rgb=c; p.alignment=a; p.space_after=Pt(5)
    return sh

def pill(s,t,x,y,w,fill=LIGHT,color=NAVY):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.36)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=t; p.alignment=PP_ALIGN.CENTER; p.font.name=FONT; p.font.size=Pt(10); p.font.bold=True; p.font.color.rgb=color

def conn(s,x1,y1,x2,y2,color=BLUE,width=1.5):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); c.line.color.rgb=color; c.line.width=Pt(width)

# 1
s=prs.slides.add_slide(blank); bg(s); rail=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(.22),Inches(7.5)); rail.fill.solid(); rail.fill.fore_color.rgb=TEAL; rail.line.fill.background()
text(s,'L0系④',.65,.52,2,.42,15,TEAL,True); text(s,'誘導集電・非接触電力伝送',.65,1.04,11.3,.72,30,NAVY,True); text(s,'相互インダクタンス → 結合回路 → 電力・効率',.67,1.82,10.8,.48,18,GRAY); pill(s,'第二種電験｜一次＋二次',.67,2.48,2.35,LIGHT,TEAL); pill(s,'Topic 30',3.16,2.48,1.18)
for i,(a,b) in enumerate([('地上側コイル','電流 i₁'),('磁界','相互磁束'),('車上側コイル','誘導起電力 e₂'),('負荷','P・η・力率')]):
    x=.85+i*2.75; card(s,[(a,15,True,NAVY,PP_ALIGN.CENTER),(b,11,False,GRAY,PP_ALIGN.CENTER)],x,3.32,2.15,1.15)
    if i<3: text(s,'→',x+2.26,3.62,.55,.4,21,TEAL,True,PP_ALIGN.CENTER)
card(s,[('M = k√(L₁L₂)',20,True,NAVY,PP_ALIGN.CENTER),('|E₂| = ωM|I₁|',19,True,TEAL,PP_ALIGN.CENTER)],.85,4.86,5.5,1.26,LIGHT2,LIGHT2)
card(s,[('実車境界',12,True,ORANGE,PP_ALIGN.LEFT),('公開一次資料で確認した原理・採用事実だけを実車事実として扱う。',14,False,NAVY,PP_ALIGN.LEFT),('k・L・M・R・f・電力・効率などの数値例は教材用仮定値。',13,False,GRAY,PP_ALIGN.LEFT)],6.65,4.86,5.5,1.26)
text(s,'狙い：非接触給電を「磁気結合された交流回路」として、二種の式で解ける形にする。',.85,6.55,11.35,.38,14,NAVY,True)
# 2
s=prs.slides.add_slide(blank); bg(s); title(s,'1｜相互誘導と結合係数','BASICS')
card(s,[('コイル1',14,True,NAVY,PP_ALIGN.CENTER),('L₁, i₁',12,False,GRAY,PP_ALIGN.CENTER)],.72,1.22,2,.86); card(s,[('コイル2',14,True,NAVY,PP_ALIGN.CENTER),('L₂, e₂',12,False,GRAY,PP_ALIGN.CENTER)],4.25,1.22,2,.86); conn(s,2.72,1.65,4.18,1.65,TEAL,2.5); text(s,'相互磁束 Ψ₁₂',2.88,1.26,1.15,.28,11,TEAL,True,PP_ALIGN.CENTER); text(s,'M',3.31,1.73,.32,.28,12,TEAL,True,PP_ALIGN.CENTER); text(s,'磁気回路の橋： Φ = NI/ℜ → L = NΦ/I。磁束分配から L₁・L₂・M を読む。',.82,2.08,5.35,.28,10.5,GRAY)
card(s,[('定義',11,True,TEAL,PP_ALIGN.LEFT),('Ψ₂ = M i₁ + L₂ i₂',19,True,NAVY,PP_ALIGN.LEFT),('e₂ = −d(Mi₁)/dt',18,True,NAVY,PP_ALIGN.LEFT),('M一定なら  e₂ = −M di₁/dt',14,False,GRAY,PP_ALIGN.LEFT)],.72,2.32,5.54,2.15)
card(s,[('結合係数',11,True,TEAL,PP_ALIGN.LEFT),('k = M / √(L₁L₂)',20,True,NAVY,PP_ALIGN.LEFT),('M² ≤ L₁L₂ ⇒ 0 ≤ |k| ≤ 1',17,True,NAVY,PP_ALIGN.LEFT),('磁気エネルギーの正半定値条件が上限を与える。',13,False,GRAY,PP_ALIGN.LEFT)],6.58,1.22,5.95,2.18)
card(s,[('正弦定常状態',11,True,TEAL,PP_ALIGN.LEFT),('|E₂| = ωM|I₁| = 2πfM|I₁|',20,True,NAVY,PP_ALIGN.LEFT),('単位確認：s⁻¹ × H × A = V',13,False,GRAY,PP_ALIGN.LEFT)],6.58,3.65,5.95,1.42,LIGHT2,LIGHT2)
card(s,[('教材用仮定値の確認例',12,True,ORANGE,PP_ALIGN.LEFT),('L₁=8.0 mH, L₂=18 mH, M=6.0 mH → k=0.500',14,False,NAVY,PP_ALIGN.LEFT),('f=50 Hz, I₁=2.0 A rms → |E₂|=3.77 V',14,False,NAVY,PP_ALIGN.LEFT)],.72,4.82,5.54,1.42,LIGHT2,LIGHT2)
card(s,[('試験への接続',12,True,TEAL,PP_ALIGN.LEFT),('R7一次：相互インダクタンス・磁気エネルギー・誘導起電力（5要素）',13,False,NAVY,PP_ALIGN.LEFT),('H30一次理論：磁気回路から L₁, L₂, M を読む（5要素）',13,False,NAVY,PP_ALIGN.LEFT)],6.58,5.33,5.95,1.32)
# 3
s=prs.slides.add_slide(blank); bg(s); title(s,'2｜結合回路を「電力」まで解く','CIRCUIT')
card(s,[('一次側',12,True,TEAL,PP_ALIGN.LEFT),('V₁, I₁',17,True,NAVY,PP_ALIGN.LEFT),('R₁ + jωL₁',13,False,GRAY,PP_ALIGN.LEFT)],.72,1.18,2.23,1.18); text(s,'jωM',3.15,1.54,1,.34,18,TEAL,True,PP_ALIGN.CENTER); conn(s,2.96,1.76,4.28,1.76,TEAL,2.2); card(s,[('二次側',12,True,TEAL,PP_ALIGN.LEFT),('I₂',17,True,NAVY,PP_ALIGN.LEFT),('R₂ + R_L + jX₂',13,False,GRAY,PP_ALIGN.LEFT)],4.4,1.18,2.33,1.18); card(s,[('負荷',12,True,ORANGE,PP_ALIGN.LEFT),('P_L = |I₂|²R_L',17,True,NAVY,PP_ALIGN.LEFT),('η = P_out/P_in',14,False,GRAY,PP_ALIGN.LEFT)],7.05,1.18,2.6,1.18,LIGHT2,LIGHT2); card(s,[('交流電力',12,True,TEAL,PP_ALIGN.LEFT),('P = VI cosφ',17,True,NAVY,PP_ALIGN.LEFT),('S = VI, Q = VI sinφ',13,False,GRAY,PP_ALIGN.LEFT)],9.95,1.18,2.55,1.18)
card(s,[('複素数で立式',12,True,TEAL,PP_ALIGN.LEFT),('V₁ = (R₁+jωL₁)I₁ + jωMI₂',16,True,NAVY,PP_ALIGN.LEFT),('0 = Z₂I₂ + jωMI₁ → I₂ = −jωMI₁/Z₂',15,False,NAVY,PP_ALIGN.LEFT),('Z_ref = (ωM)²/Z₂',15,True,BLUE,PP_ALIGN.LEFT),('力率は入力側の V₁ と I₁ の位相差から読む。',13,False,GRAY,PP_ALIGN.LEFT)],.72,2.67,5.86,2.36)
card(s,[('教材用仮定値｜本試験標準例',12,True,ORANGE,PP_ALIGN.LEFT),('L₁=2.0 mH, L₂=1.5 mH, k=0.30, f=1 kHz',13,False,NAVY,PP_ALIGN.LEFT),('R₁=0.40 Ω, R₂=0.30 Ω, R_L=4.0 Ω, I₁=10.0 A',13,False,NAVY,PP_ALIGN.LEFT),('M=0.5196 mH / C₂=16.8869 μF',14,True,NAVY,PP_ALIGN.LEFT),('I₂=7.593 A → P_L=230.59 W',14,True,NAVY,PP_ALIGN.LEFT),('P_cu1=40.00 W, P_cu2=17.29 W → η=80.1%',14,True,TEAL,PP_ALIGN.LEFT)],6.89,2.67,5.61,2.8,LIGHT2,LIGHT2)
for i,t in enumerate(['① Mを求める','② Z₂を作る','③ I₂を求める','④ P・力率を計算','⑤ 損失→η']):
    x=.72+i*2.41; pill(s,t,x,5.72,2.15,WHITE,NAVY)
    if i<4: text(s,'→',x+2.15,5.73,.27,.30,14,TEAL,True,PP_ALIGN.CENTER)
text(s,'「磁気」だけで終わらず、最終的に交流電力・損失・効率へ落とす。',.78,6.39,11.8,.36,13,GRAY)
# 4
s=prs.slides.add_slide(blank); bg(s); title(s,'3｜3つの可視化で挙動をつかむ','VISUAL'); pill(s,'すべて教材用仮定値。L0系実車性能ではない。',8.65,.36,3.85,RGBColor(252,238,224),ORANGE)
xs=[.48,4.45,8.42]
for i,p in enumerate(charts):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(xs[i]),Inches(1.05),Inches(3.93),Inches(4.85)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.color.rgb=LIGHT; s.shapes.add_picture(p,Inches(xs[i]+.12),Inches(1.20),width=Inches(3.69))
for i,(a,b,c) in enumerate([('固定 I₁・共振条件では','P_L ∝ k²','kを2倍 → 電力は約4倍。'),('負荷を変えると','η に山が現れる','小さすぎても大きすぎても損失比率が悪化。'),('リアクタンスが相殺する','f₀ 付近で応答最大','周波数ずれで |Z| が増え、電流・電力が低下。')]):
    text(s,a,xs[i]+.20,4.33,3.48,.27,11,GRAY); text(s,b,xs[i]+.20,4.64,3.48,.33,15,TEAL,True); text(s,c,xs[i]+.20,5.05,3.48,.55,11,NAVY)
card(s,[('読み方',12,True,TEAL,PP_ALIGN.LEFT),('k・負荷・f のどれを動かしたグラフか、他の条件は何を固定したかを先に確認する。',14,False,NAVY,PP_ALIGN.LEFT)],.72,6.15,11.9,.72,LIGHT2,LIGHT2)
# 5
s=prs.slides.add_slide(blank); bg(s); title(s,'4｜共振・損失・最大効率を二種水準へ','EXAM BRIDGE')
card(s,[('RLC直列共振',14,True,TEAL,PP_ALIGN.LEFT),('Z = R + j(ωL − 1/ωC)',18,True,NAVY,PP_ALIGN.LEFT),('ω₀ = 1/√(LC)',18,True,NAVY,PP_ALIGN.LEFT),('共振時：I₀=V/R',15,False,NAVY,PP_ALIGN.LEFT),('Q=ω₀L/R=1/(ω₀CR)',15,False,NAVY,PP_ALIGN.LEFT),('V_L = V_C = QV',15,False,NAVY,PP_ALIGN.LEFT),('→ 周波数特性・力率1の基礎',13,False,GRAY,PP_ALIGN.LEFT)],.72,1.18,3.73,3.87)
card(s,[('損失と変換効率',14,True,TEAL,PP_ALIGN.LEFT),('P_in = P_out + Σ損失',18,True,NAVY,PP_ALIGN.LEFT),('η = P_out / P_in',18,True,NAVY,PP_ALIGN.LEFT),('変圧器橋渡し：',13,True,ORANGE,PP_ALIGN.LEFT),('P_c(α)=α²P_c,n',15,False,NAVY,PP_ALIGN.LEFT),('最大効率：P_i = α²P_c,n',15,True,NAVY,PP_ALIGN.LEFT),('α_max = √(P_i/P_c,n)',15,True,NAVY,PP_ALIGN.LEFT)],4.78,1.18,3.76,3.87); text(s,'同一磁束密度：ヒステリシス損 ∝ f、渦電流損 ∝ f²',4.95,4.48,3.40,.42,10,GRAY)
card(s,[('教材用仮定値｜複合例',13,True,ORANGE,PP_ALIGN.LEFT),('S_n=100 kVA',14,False,NAVY,PP_ALIGN.LEFT),('P_i=500 W, P_c,n=800 W',14,False,NAVY,PP_ALIGN.LEFT),('α_max = 0.7906',17,True,NAVY,PP_ALIGN.LEFT),('α=0.60, cosφ=0.80',14,False,NAVY,PP_ALIGN.LEFT),('P_out=48.0 kW',15,True,NAVY,PP_ALIGN.LEFT),('η = 98.4%',18,True,TEAL,PP_ALIGN.LEFT)],8.87,1.18,3.66,3.87,LIGHT2,LIGHT2); text(s,'η_max ≈ 98.75%（cosφ=1）',9.04,4.62,3.25,.30,10.5,TEAL,True)
card(s,[('試験への接続',12,True,TEAL,PP_ALIGN.LEFT),('R8一次：RLC共振・Q・端子電圧（5要素） ｜ H30一次機械：鉄損・銅損・試験・規約効率（5要素）',13,False,NAVY,PP_ALIGN.LEFT),('R2二次：無負荷損・全負荷銅損・最大効率条件・負荷率/力率/効率の記述計算（5要素）',13,False,NAVY,PP_ALIGN.LEFT)],.72,5.36,11.81,1.18)
text(s,'変圧器問題は「結合回路・電力・損失・効率」の橋渡し。L0系実機と同一構造・同一定数とはみなさない。',.77,6.73,11.7,.34,11,RED,True)
# 6
s=prs.slides.add_slide(blank); bg(s); title(s,'5｜固定過去問への接続と解法フロー','GATE'); text(s,'固定過去問',.72,1.18,2.72,.30,11,GRAY,True); text(s,'教材で接続する5答案要素',3.62,1.18,7.55,.30,11,GRAY,True); text(s,'要素',11.47,1.18,.70,.30,11,GRAY,True,PP_ALIGN.CENTER)
rows=[('R8 一次 理論 問5','RLC共振 / I₀ / Q / L・C端子電圧 / 共振条件'),('R7 一次 理論 問2','M / 磁気エネルギー / 誘導起電力 / M²≤L₁L₂ / k'),('H30 一次 理論 問2','磁気回路 / 磁束分配 / L₁ / L₂ / M'),('H30 一次 機械 問5','無負荷・短絡試験 / 鉄損(f・f²) / 銅損 / 規約効率 / 周波数依存'),('R2 二次 機械・制御 問2','無負荷損 / 全負荷銅損 / α_max / η_max / 指定負荷率・力率でのη')]
for i,(exam,elems) in enumerate(rows):
    y=1.57+i*.70; sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.68),Inches(y),Inches(11.72),Inches(.58)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE if i%2==0 else LIGHT2; sh.line.fill.background(); text(s,exam,.82,y+.11,2.65,.28,12,NAVY,True); text(s,elems,3.62,y+.11,7.52,.28,11,NAVY); pill(s,'5',11.58,y+.10,.50,LIGHT,TEAL)
text(s,'一次 20/20 ＋ 二次 5/5 ＝ 25/25 connected',8.05,5.16,4.28,.35,13,TEAL,True,PP_ALIGN.RIGHT)
card(s,[('解法フロー',12,True,TEAL,PP_ALIGN.LEFT),('① 磁気回路・L・M・k → ② 複素等価回路 → ③ 共振・I → ④ P・力率 → ⑤ 損失・η',15,True,NAVY,PP_ALIGN.LEFT)],.72,5.55,11.72,.84)
card(s,[('境界',11,True,ORANGE,PP_ALIGN.LEFT),('未確認のL0系実車定数は真値化しない。Topic 31以降は先取りしない。Topic 21の一般式は変更しない。',12,False,NAVY,PP_ALIGN.LEFT)],.72,6.54,11.72,.64,RGBColor(252,244,234),RGBColor(242,215,188))

prs.core_properties.title='L0系④ 誘導集電・非接触電力伝送'; prs.core_properties.subject='電験二種 Topic 30 解説画像'; prs.core_properties.author='cureflash / OpenAI worker'
path=os.path.join(OUT,'30_l0_inductive_power_collection_contactless_power_transfer_images.pptx'); prs.save(path); print(path)
