from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
import math, os, zipfile, hashlib

OUT='denken-shinkansen/10_electrical_math_2/topics/17_ode_numerical_methods/17_ode_numerical_methods_images.pptx'
prs=Presentation()
prs.slide_width=Inches(13.333)
prs.slide_height=Inches(7.5)
FONT='Noto Sans CJK JP'
MONO='Noto Sans Mono CJK JP'
BG=RGBColor(247,249,252); NAVY=RGBColor(26,43,74); BLUE=RGBColor(47,111,178); TEAL=RGBColor(35,145,141)
ORANGE=RGBColor(222,132,59); RED=RGBColor(186,69,69); GRAY=RGBColor(92,103,116); LGRAY=RGBColor(223,228,235)
WHITE=RGBColor(255,255,255); DARK=RGBColor(32,37,45); GREEN=RGBColor(64,140,92)

def add_bg(slide):
    bg=slide.background; bg.fill.solid(); bg.fill.fore_color.rgb=BG
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(0.18)); bar.fill.solid(); bar.fill.fore_color.rgb=NAVY; bar.line.fill.background()

def add_title(slide,title,subtitle=None):
    tb=slide.shapes.add_textbox(Inches(0.55),Inches(0.28),Inches(12.2),Inches(0.55)); p=tb.text_frame.paragraphs[0]
    p.text=title; p.font.name=FONT; p.font.size=Pt(27); p.font.bold=True; p.font.color.rgb=NAVY
    if subtitle:
        sb=slide.shapes.add_textbox(Inches(0.58),Inches(0.83),Inches(12.0),Inches(0.38)); p=sb.text_frame.paragraphs[0]
        p.text=subtitle; p.font.name=FONT; p.font.size=Pt(12); p.font.color.rgb=GRAY

def add_card(slide,x,y,w,h,title=None,fill=WHITE):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=LGRAY; sh.line.width=Pt(1)
    if title:
        tb=slide.shapes.add_textbox(Inches(x+0.18),Inches(y+0.12),Inches(w-0.36),Inches(0.38)); p=tb.text_frame.paragraphs[0]
        p.text=title; p.font.name=FONT; p.font.size=Pt(17); p.font.bold=True; p.font.color.rgb=NAVY
    return sh

def add_text(slide,x,y,w,h,text,size=16,color=DARK,bold=False,align=PP_ALIGN.LEFT,font=FONT,margin=0.05):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(margin); p=tf.paragraphs[0]
    p.text=text; p.alignment=align; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color; return tb

def add_bullets(slide,x,y,w,h,items,size=15,color=DARK):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.06); tf.margin_top=tf.margin_bottom=Inches(0.02)
    for idx,item in enumerate(items):
        p=tf.paragraphs[0] if idx==0 else tf.add_paragraph(); p.text='• '+item; p.font.name=FONT; p.font.size=Pt(size); p.font.color.rgb=color; p.space_after=Pt(7)
    return tb

def add_formula(slide,x,y,w,h,text,size=20,fill=RGBColor(240,245,250)):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=RGBColor(187,201,217)
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]
    p.text=text; p.alignment=PP_ALIGN.CENTER; p.font.name=MONO; p.font.size=Pt(size); p.font.bold=True; p.font.color.rgb=NAVY; return sh

def add_footer(slide,n): add_text(slide,11.75,7.03,1.0,0.25,f'Topic 17  |  {n}/5',size=9,color=GRAY,align=PP_ALIGN.RIGHT)

# 1: IVP and Euler
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s); add_title(s,'常微分方程式の数値解法','初期値問題を「時間刻みで追跡する」ための最小セット')
add_card(s,0.6,1.35,4.0,4.9,'1. 初期値問題'); add_formula(s,0.95,2.0,3.3,0.95,'dy/dt = f(t, y)\ny(t₀) = y₀',22)
add_bullets(s,0.95,3.15,3.25,2.3,['微分方程式だけでなく初期値が必要','回路なら KCL/KVL から dy/dt の形へ','列車なら運動方程式から dv/dt の形へ','単位・符号・定常値を先に確認'],14)
add_card(s,4.85,1.35,7.85,4.9,'2. Euler法の更新'); add_formula(s,5.25,2.0,7.0,0.72,'tₙ₊₁ = tₙ + h',23); add_formula(s,5.25,2.92,7.0,0.85,'yₙ₊₁ = yₙ + h f(tₙ, yₙ)',23)
labels=['式を整理','初期値','刻み幅 h','傾き f','1歩更新','検算']; xs=[5.15,6.35,7.55,8.75,9.95,11.15]
for i,(x,lab) in enumerate(zip(xs,labels)):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(4.15),Inches(1.05),Inches(0.75)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(234,240,247) if i<5 else RGBColor(233,245,238); sh.line.color.rgb=BLUE if i<5 else GREEN
    tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=lab; p.alignment=PP_ALIGN.CENTER; p.font.name=FONT; p.font.size=Pt(12); p.font.bold=True; p.font.color.rgb=NAVY
    if i<5:
        a=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+0.92),Inches(4.34),Inches(0.35),Inches(0.33)); a.fill.solid(); a.fill.fore_color.rgb=RGBColor(164,177,193); a.line.fill.background()
add_text(s,5.25,5.28,7.05,0.62,'検算は「刻み幅変更・解析解・定常値」のいずれかで行う。',16,RED,True,PP_ALIGN.CENTER); add_footer(s,1)

# 2: step size and error
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s); add_title(s,'Euler法と刻み幅','dy/dt = -y,  y(0)=1 で誤差の挙動を見る')
add_card(s,0.6,1.28,5.0,5.35,'更新式'); add_formula(s,0.95,1.95,4.3,0.72,'yₙ₊₁ = (1-h)yₙ',21); add_text(s,0.95,2.95,4.25,0.45,'h = 0.25 の4ステップ',16,NAVY,True)
vals=['1.000000','0.750000','0.562500','0.421875','0.316406']
for i,v in enumerate(vals):
    x=0.95+i*0.84; sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(3.55),Inches(0.68),Inches(0.64)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(239,245,252); sh.line.color.rgb=BLUE
    tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=v; p.font.name=MONO; p.font.size=Pt(10); p.alignment=PP_ALIGN.CENTER; p.font.color.rgb=DARK
    if i<4:
        a=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+0.62),Inches(3.72),Inches(0.24),Inches(0.25)); a.fill.solid(); a.fill.fore_color.rgb=RGBColor(164,177,193); a.line.fill.background()
add_formula(s,0.95,4.55,4.3,0.8,'解析解:  y(1)=e⁻¹=0.367879...',18,RGBColor(247,241,232)); add_text(s,0.95,5.55,4.25,0.6,'絶対誤差 = |0.31640625 - 0.367879...|\n= 0.051473...',14,RED,True,PP_ALIGN.CENTER)
add_card(s,5.85,1.28,6.85,5.35,'刻み幅を変えて比較'); headers=['h','Euler y(1)','絶対誤差']; data=[('0.5000','0.250000','0.117879'),('0.2500','0.316406','0.051473'),('0.1250','0.343609','0.024271'),('0.0625','0.356074','0.011805')]; colx=[6.2,8.15,10.35]; colw=[1.6,1.9,1.9]
for x,w,hdr in zip(colx,colw,headers):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(2.0),Inches(w),Inches(0.5)); sh.fill.solid(); sh.fill.fore_color.rgb=NAVY; sh.line.color.rgb=WHITE; tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=hdr; p.alignment=PP_ALIGN.CENTER; p.font.name=FONT; p.font.size=Pt(13); p.font.bold=True; p.font.color.rgb=WHITE
for r,row in enumerate(data):
    y=2.52+r*0.65
    for x,w,val in zip(colx,colw,row):
        sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(0.6)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE if r%2==0 else RGBColor(245,247,250); sh.line.color.rgb=LGRAY; tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=val; p.alignment=PP_ALIGN.CENTER; p.font.name=MONO; p.font.size=Pt(13); p.font.color.rgb=DARK
add_text(s,6.18,5.42,6.0,0.42,'この例では h を小さくすると誤差が単調に減少',15,GREEN,True,PP_ALIGN.CENTER)
errs=[0.117879,0.051473,0.024271,0.011805]
for i,e in enumerate(errs):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(6.6),Inches(5.92+i*0.16),Inches(5.0*(e/max(errs))),Inches(0.09)); sh.fill.solid(); sh.fill.fore_color.rgb=TEAL; sh.line.fill.background()
add_footer(s,2)

# 3: RL transient
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s); add_title(s,'RL過渡をEuler法で追う','解析解を知る回路で、数値解の更新と誤差検算を同時に練習')
add_card(s,0.6,1.28,4.2,5.35,'回路式 → 初期値問題'); add_formula(s,0.9,1.9,3.6,0.82,'L di/dt + Ri = V',19); add_formula(s,0.9,2.92,3.6,0.82,'di/dt = (V-Ri)/L',19)
add_bullets(s,0.95,3.95,3.45,1.35,['V=100 V, R=10 Ω, L=2 H','i(0)=0 A, 定常値=10 A','時定数 τ=L/R=0.20 s'],14); add_formula(s,0.9,5.45,3.6,0.72,'解析解: i=10(1-e⁻⁵ᵗ)',16,RGBColor(247,241,232))
add_card(s,5.05,1.28,7.65,5.35,'t = 0.40 s まで計算'); add_text(s,5.38,1.93,3.1,0.35,'h = 0.10 s',16,BLUE,True,PP_ALIGN.CENTER); add_text(s,9.0,1.93,3.1,0.35,'h = 0.05 s',16,TEAL,True,PP_ALIGN.CENTER)
add_formula(s,5.38,2.38,3.1,0.65,'0 → 5 → 7.5 → 8.75 → 9.375 A',13); add_formula(s,9.0,2.38,3.1,0.65,'8刻み → 8.998871 A',14); add_text(s,5.4,3.34,3.1,0.42,'絶対誤差 0.728353 A',14,RED,True,PP_ALIGN.CENTER); add_text(s,9.0,3.34,3.1,0.42,'絶対誤差 0.352224 A',14,GREEN,True,PP_ALIGN.CENTER); add_formula(s,6.15,4.12,5.45,0.8,'解析解 i(0.40)=8.646647 A',19,RGBColor(247,241,232))
x0,y0,w,h=5.55,5.15,6.3,0.85
for xx,yy,ww,hh in [(x0,y0+h,w,0.015),(x0,y0,0.015,h)]:
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(xx),Inches(yy),Inches(ww),Inches(hh)); sh.fill.solid(); sh.fill.fore_color.rgb=GRAY; sh.line.fill.background()
def px(t): return x0+w*(t/.4)
def py(i): return y0+h*(1-i/10)
pts=[(px(.05*k),py(10*(1-math.exp(-5*.05*k)))) for k in range(9)]
for (xa,ya),(xb,yb) in zip(pts,pts[1:]):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(xa),Inches(ya),Inches(xb),Inches(yb)); ln.line.color.rgb=ORANGE; ln.line.width=Pt(2)
epts=[(px(.1*k),py(v)) for k,v in enumerate([0,5,7.5,8.75,9.375])]
for (xa,ya),(xb,yb) in zip(epts,epts[1:]):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(xa),Inches(ya),Inches(xb),Inches(yb)); ln.line.color.rgb=BLUE; ln.line.width=Pt(2)
add_text(s,5.65,6.15,6.1,0.32,'橙: 解析解   青: Euler (h=0.10 s)',11,GRAY,False,PP_ALIGN.CENTER); add_footer(s,3)

# 4: train IVP
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s); add_title(s,'簡略列車モデルを時間発展させる','実車値ではなく教育用の仮定モデルとして、速度と位置を同時更新')
add_card(s,0.6,1.28,5.2,5.35,'運動方程式'); add_formula(s,0.95,1.92,4.5,0.78,'m dv/dt = Ftraction(v) - Fresistance(v)',17); add_formula(s,0.95,2.88,4.5,0.78,'dv/dt = 0.5 - 0.0025v',18); add_formula(s,0.95,3.84,4.5,0.78,'dx/dt = v',20)
add_bullets(s,0.95,4.92,4.5,1.18,['h=20 s, v(0)=0, x(0)=0','各刻みの始点値 vₙ で速度・位置を更新','係数は教材用仮定値。実車値ではない'],13)
add_card(s,6.05,1.28,6.65,5.35,'Euler更新結果'); rows=[('0 s','0.000 m/s','0 m'),('20 s','10.000 m/s','0 m'),('40 s','19.500 m/s','200 m'),('60 s','28.525 m/s','590 m')]; headers=['時刻','速度 v','位置 x']; cols=[6.35,8.15,10.25]; widths=[1.55,1.85,1.85]
for x,w,hdr in zip(cols,widths,headers):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(1.96),Inches(w),Inches(0.5)); sh.fill.solid(); sh.fill.fore_color.rgb=NAVY; sh.line.color.rgb=WHITE; tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=hdr; p.alignment=PP_ALIGN.CENTER; p.font.name=FONT; p.font.size=Pt(13); p.font.bold=True; p.font.color.rgb=WHITE
for r,row in enumerate(rows):
    y=2.48+r*0.66
    for x,w,val in zip(cols,widths,row):
        sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(0.60)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE if r%2==0 else RGBColor(245,247,250); sh.line.color.rgb=LGRAY; tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=val; p.alignment=PP_ALIGN.CENTER; p.font.name=MONO; p.font.size=Pt(12); p.font.color.rgb=DARK
add_formula(s,6.42,5.35,5.75,0.72,'vₙ₊₁ = vₙ + 20(0.5 - 0.0025vₙ)',15); add_formula(s,6.42,6.14,5.75,0.55,'xₙ₊₁ = xₙ + 20vₙ',15); add_footer(s,4)

# 5: RK concept and exam boundary
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s); add_title(s,'Runge-Kutta法の概念と試験への接続','本Topicでは概念まで。確認済み第二種過去問は解析的な過渡問題への隣接接続として扱う')
add_card(s,0.6,1.28,5.35,5.35,'Runge-Kutta法: 1刻みで複数の傾きを見る'); x0=1.05; y0=4.8
for xx,yy,ww,hh in [(x0,y0,4.3,0.015),(x0,2.0,0.015,2.8)]:
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(xx),Inches(yy),Inches(ww),Inches(hh)); sh.fill.solid(); sh.fill.fore_color.rgb=GRAY; sh.line.fill.background()
curve=[(1.05,4.3),(1.8,3.95),(2.55,3.55),(3.3,3.2),(4.05,2.92),(4.8,2.73)]
for (xa,ya),(xb,yb) in zip(curve,curve[1:]):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(xa),Inches(ya),Inches(xb),Inches(yb)); ln.line.color.rgb=TEAL; ln.line.width=Pt(2.4)
for x,y,label in [(1.35,4.16,'k₁'),(2.55,3.55,'k₂'),(3.3,3.2,'k₃'),(4.45,2.80,'k₄')]:
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x-0.28),Inches(y+0.12),Inches(x+0.28),Inches(y-0.12)); ln.line.color.rgb=ORANGE; ln.line.width=Pt(2); add_text(s,x-0.18,y-0.35,0.5,0.28,label,11,ORANGE,True,PP_ALIGN.CENTER)
add_text(s,1.0,5.2,4.55,0.9,'Euler法は始点の傾き1個。\nRunge-Kutta法は刻み途中を含む複数の傾きを組み合わせる。',15,DARK,False,PP_ALIGN.CENTER); add_text(s,1.0,6.22,4.55,0.34,'次数一般論・適応刻み・多段法には広げない。',13,RED,True,PP_ALIGN.CENTER)
add_card(s,6.18,1.28,6.52,5.35,'第二種過去問への接続と境界'); add_text(s,6.52,1.88,5.8,0.42,'接続確認した一次「理論」5問',16,NAVY,True)
add_bullets(s,6.52,2.34,5.8,2.3,['R8 問6 (1)〜(2): 一次応答の微分方程式・初期値','R7 問6 / R6 問4: 過渡応答・定常値・時定数','R5 問5 / R4 問4: RL/RC方程式と初期条件','直接Euler法・Runge-Kutta法を要求する確認済み問題: 0問','この5問を固定品質ゲート件数へ水増ししない'],13)
add_text(s,6.52,4.88,5.8,0.35,'頻出ミス',15,RED,True); add_bullets(s,6.52,5.25,5.8,1.12,['初期値を入れない','hの単位をそろえない','更新時点を混在させる','教育用列車係数を実車値と扱う'],12); add_footer(s,5)

prs.save(OUT)
with zipfile.ZipFile(OUT,'r') as z:
    if z.testzip() is not None: raise SystemExit('PPTX ZIP validation failed')
with open(OUT,'rb') as f: data=f.read()
print('PPTX',len(data),'bytes','sha256',hashlib.sha256(data).hexdigest())
print('slides',len(prs.slides),'size',prs.slide_width,prs.slide_height)
if len(prs.slides)!=5: raise SystemExit('slide count mismatch')
