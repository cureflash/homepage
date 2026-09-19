from pathlib import Path
from hashlib import sha256
from pptx import Presentation

p = Path('denken-shinkansen/05_shinkansen_vehicle_2/topics/25_n700s_battery_self_propulsion_energy_design/25_n700s_battery_self_propulsion_energy_design_images.pptx')
metrics = Path('denken-shinkansen/05_shinkansen_vehicle_2/topics/25_n700s_battery_self_propulsion_energy_design/_tmp_topic25_ppt_metrics.txt')
prs = Presentation(p)
assert len(prs.slides) == 6
sw, sh = prs.slide_width, prs.slide_height
texts = []
for idx, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (idx, shape.name)
        assert shape.left + shape.width <= sw + 1, (idx, shape.name)
        assert shape.top + shape.height <= sh + 1, (idx, shape.name)
        if hasattr(shape, 'text'):
            texts.append(shape.text)
all_text = '\n'.join(texts)
for token in [
    'SCiB', 'Ah', 'Wh', '直列', '並列', 'Cレート', '放電電流', '電池効率',
    'インバータ効率', 'モーター効率', '必要容量', 'PbO₂', 'PbSO₄',
    '26.8 Ah/mol', '3.7 V', 'Vterm = Ei + IRi', 'Vterm = Ei − IRi',
    'Eb1 = D Ep1', 'Ep2 = Eb2/(1−D)', '走行時間―必要電力量',
    '速度条件―必要電力', '効率―必要電池容量', '28答案要素', 'H23 二次'
]:
    assert token in all_text, token
raw = p.read_bytes()
metrics.write_text(f'size={len(raw)}\nsha256={sha256(raw).hexdigest()}\nslides=6\n', encoding='utf-8')
print('package/open/geometry/content PASS, slides=6')
